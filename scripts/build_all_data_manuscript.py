"""Build a replayable, evidence-bounded manuscript from the elastomer project.

The script is intentionally read-only with respect to the project data and the
immutable evidence package.  It writes a new versioned manuscript directory,
an exhaustive source/run audit, a claim--evidence--literature matrix, and (when
``python-docx`` is available) an editable Word document.

The external project is kept outside the repository, so this module accepts an
explicit project root.  No scientific value is recomputed here: approved CSV
tables are read and rendered, while package run manifests are used only for
provenance and quality boundaries.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable, Mapping


# Make the script usable both as ``python scripts/...`` and as an imported test
# module without requiring the caller to install the package first.
REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

SAMPLES: tuple[str, ...] = ("PA6", "PA6-50", "PA11", "PA11-50", "PA12", "PA12-50")
TECHNIQUES: tuple[str, ...] = (
    "DSC",
    "IR",
    "NMR-solution",
    "NMR-solid",
    "WAXS",
    "SAXS",
)

RAW_SPECS: tuple[dict[str, Any], ...] = (
    {
        "root": "DSC-等温结晶",
        "extensions": (".txt",),
        "technique": "DSC",
        "mode": "isothermal_dsc",
    },
    {
        "root": "DSC-升降升",
        "extensions": (".txt", ".opju"),
        "technique": "DSC",
        "mode": "nonisothermal_dsc",
    },
    {
        "root": "insu-FTIR",
        "extensions": (".csv", ".spc"),
        "technique": "IR",
        "mode": "temperature_series",
    },
    {
        "root": "nmr",
        "extensions": (".csv",),
        # Keep the public technique name aligned with the shared PolyNexus
        # contract.  The acquisition mode carries the solution/solid split.
        "technique": "NMR",
        "mode": "solution_1d",
    },
    {
        "root": "solid-NMR",
        "extensions": (".csv",),
        "technique": "NMR",
        "mode": "solid_1d",
    },
    {
        "root": "waxs",
        "extensions": (".raw",),
        "technique": "WAXS",
        "mode": "profile_from_detector",
    },
    {
        "root": "saxs",
        "extensions": (".edf",),
        "technique": "SAXS",
        "mode": "profile_from_detector",
    },
)

DERIVED_TABLES: tuple[tuple[str, str, str], ...] = (
    (
        "isothermal_dsc_kinetics",
        "analysis_output/control_dsc/dsc_kinetics.csv",
        "primary_kinetics",
    ),
    (
        "isothermal_dsc_kinetics_pa50",
        "analysis_output/paper_data/dsc_kinetics.csv",
        "primary_kinetics_pa50",
    ),
    (
        "isothermal_dsc_pairwise",
        "analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv",
        "primary_pairwise_comparison",
    ),
    (
        "isothermal_dsc_qc",
        "analysis_output/control_dsc/dsc_qc_review.csv",
        "quality_audit",
    ),
    (
        "nonisothermal_dsc_summary",
        "analysis_output/paper_figures_pairwise/nonisothermal_dsc_thermal_summary.csv",
        "specified_thermal_context",
    ),
    (
        "ftir_2dcos_summary",
        "analysis_output/paper_data/ftir_2dcos_pairwise_summary.csv",
        "specialized_pairwise_support",
    ),
    (
        "scattering_profile_summary",
        "analysis_output/paper_data/scattering_profile_summary.csv",
        "qualitative_scattering_index",
    ),
    (
        "scattering_profiles",
        "analysis_output/paper_data/scattering_profiles.csv",
        "normalized_profile_data",
    ),
)


def _json_load(path: Path, default: Any = None) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError):
        return default


def _norm_path(value: str | Path) -> str:
    """Normalize a path for suffix matching across Windows/JSON exports."""

    text = str(value).replace("\\", "/").strip().lower()
    # Drive letters and duplicated separators are irrelevant for provenance
    # matching; retain non-ASCII directory names because they are meaningful.
    text = re.sub(r"/+", "/", text)
    return text.rstrip("/")


def _basename(value: str | Path) -> str:
    return _norm_path(value).rsplit("/", 1)[-1]


def infer_sample(value: str | Path | None) -> str:
    """Infer only the declared sample label encoded in a source path/name.

    The two numeric aliases (4012/4033) and ``buyangpin`` are existing folder
    labels in this project.  They are treated as naming aliases, not as a
    measurement of composition.
    """

    text = str(value or "")
    folded = text.replace("_", "-").replace(" ", "")
    match = re.search(r"(?i)pa(?:11|12|6)(?:-?50)?", folded)
    if match:
        token = match.group(0).upper().replace("-?", "-")
        token = token.replace("50", "-50") if token.endswith("50") and not token.endswith("-50") else token
        token = re.sub(r"PA(6|11|12)-?50$", r"PA\1-50", token)
        token = re.sub(r"PA(6|11|12)$", r"PA\1", token)
        if token in SAMPLES:
            return token
    lowered = folded.lower()
    if "4012" in lowered:
        return "PA12-50"
    if "4033" in lowered:
        return "PA11-50"
    if "buyangpin" in lowered:
        return "PA6-50"
    return "unassigned"


def _temperature_from_name(path: Path) -> str:
    stem = path.stem
    # File names begin with sample aliases such as PA11/PA12/4012 and may end
    # with a one-minute hold marker.  Read the temperature token attached to
    # the hold marker first, then the terminal token; never take the first
    # number in the name because that is commonly a sample label.
    hold = re.search(r"[-_](\d{2,3})(?=-for\b)", stem, flags=re.IGNORECASE)
    if hold:
        return hold.group(1)
    terminal = re.search(r"[-_](\d{2,3})(?:\s*°?c)?$", stem, flags=re.IGNORECASE)
    return terminal.group(1) if terminal else ""


def _export_family(path: Path, technique: str, mode: str, sample: str) -> str:
    """Return a transparent family label used for duplicate auditing."""

    stem = path.stem.lower()
    if technique == "IR":
        temperature = _temperature_from_name(path)
        if "for" in stem:
            return f"IR:{sample}:time-export:{temperature or 'unknown'}"
        return f"IR:{sample}:temperature:{temperature or 'unknown'}"
    if mode == "solid_1d":
        base = re.sub(r"txt", "", stem)
        base = re.sub(r"-70", "", base)
        return f"NMR-solid:{sample}:{base}"
    if mode == "solution_1d":
        return f"NMR-solution:{sample}:{stem}"
    if technique == "DSC":
        return f"DSC:{mode}:{sample}:{stem.lower()}"
    return f"{technique}:{sample}"


def classify_source(path: Path) -> dict[str, str]:
    """Classify one project-relative source without changing its data."""

    rel = path.as_posix()
    sample = infer_sample(rel)
    suffix = path.suffix.lower()
    technique = "other"
    mode = "unknown"
    use = "audit_only"
    reason = ""
    for spec in RAW_SPECS:
        root = spec["root"]
        if rel == root or rel.startswith(root + "/"):
            technique = str(spec["technique"])
            mode = str(spec["mode"])
            break
    if technique == "DSC" and mode == "isothermal_dsc":
        use = "primary"
    elif technique == "DSC" and mode == "nonisothermal_dsc":
        if suffix == ".opju":
            use = "audit_only"
            reason = "format_unsupported"
        else:
            use = "supporting"
    elif technique == "IR":
        use = "audit_only"
        reason = "raw_temperature_series_uses_specialized_summary"
    elif technique == "NMR":
        use = "audit_only"
        reason = "review_required_structural_inventory"
    elif technique in {"WAXS", "SAXS"}:
        use = "supporting_qualitative"
        reason = "normalized_profile_no_absolute_background"
    family = _export_family(path, technique, mode, sample)
    return {
        "technique": technique,
        "mode": mode,
        "sample": sample,
        "use": use,
        "exclusion_reason": reason,
        "export_family": family,
    }


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _iter_raw_files(project_root: Path) -> Iterable[tuple[Path, dict[str, Any]]]:
    for spec in RAW_SPECS:
        root = project_root / str(spec["root"])
        if not root.exists():
            continue
        extensions = {str(v).lower() for v in spec["extensions"]}
        for path in sorted(root.rglob("*")):
            if path.is_file() and path.suffix.lower() in extensions:
                yield path, spec


def _run_artifacts(run: Mapping[str, Any]) -> list[dict[str, Any]]:
    recipe = run.get("analysis_run", {}).get("recipe", {})
    artifacts = recipe.get("artifacts", [])
    if not artifacts:
        artifacts = run.get("recipe", {}).get("artifacts", [])
    return [dict(v) for v in artifacts if isinstance(v, Mapping)]


def _walk_dicts(value: Any) -> Iterable[dict[str, Any]]:
    if isinstance(value, Mapping):
        yield dict(value)
        for child in value.values():
            yield from _walk_dicts(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_dicts(child)


def _source_strings(run: Mapping[str, Any]) -> list[str]:
    """Collect explicit input/source references without traversing metric paths.

    Recipe artifacts are handled separately in :func:`_run_record`.  Nested
    ``path`` keys may describe result metrics rather than files, so only the
    explicit source/input keys below participate in the provenance fallback.
    """

    values: list[str] = []
    for obj in _walk_dicts(run):
        for key, value in obj.items():
            if not isinstance(value, str):
                continue
            key_lower = str(key).lower()
            if key_lower in {"source_path", "input_path"}:
                if value:
                    values.append(value)
    return list(dict.fromkeys(values))


def _first_signal(analysis: Mapping[str, Any], name: str) -> Any:
    for signal in analysis.get("confidence_signals", []) if isinstance(analysis, Mapping) else []:
        if isinstance(signal, Mapping) and signal.get("name") == name:
            return signal.get("value")
    return None


def _first_evidence_analysis(run: Mapping[str, Any]) -> dict[str, Any]:
    items = run.get("evidence_items", [])
    if not items or not isinstance(items[0], Mapping):
        return {}
    evidence = items[0].get("observed_results", {})
    if not isinstance(evidence, Mapping):
        return {}
    analysis = evidence.get("analysis_evidence", {})
    return dict(analysis) if isinstance(analysis, Mapping) else {}


def _interpretation_text(run: Mapping[str, Any]) -> str:
    values = run.get("analysis_run", {}).get("evidence", {}).get("supported_interpretations", [])
    return " | ".join(str(v) for v in values)


def _parse_numeric_from_text(text: str, key: str) -> float | None:
    match = re.search(rf"{re.escape(key)}=([-+]?\d+(?:\.\d+)?)", text)
    return float(match.group(1)) if match else None


def _run_record(run: Mapping[str, Any], ref: str, raw_paths: list[str]) -> dict[str, Any]:
    artifacts = _run_artifacts(run)
    technique = str(artifacts[0].get("technique", "unknown")) if artifacts else "unknown"
    source_role = str(run.get("request_parameters", {}).get("source_role", ""))
    if not source_role:
        source_role = str(run.get("analysis_run", {}).get("request_parameters", {}).get("source_role", ""))
    source_refs = _source_strings(run)
    top_paths = [str(v.get("path", "")) for v in artifacts if v.get("path")]
    all_refs = list(dict.fromkeys(top_paths + source_refs))
    names = [_basename(v) for v in all_refs]
    samples = sorted({infer_sample(v) for v in all_refs if infer_sample(v) != "unassigned"})
    evidence_ids = [
        str(v.get("evidence_id"))
        for v in run.get("evidence_items", [])
        if isinstance(v, Mapping) and v.get("evidence_id")
    ]
    analysis = _first_evidence_analysis(run)
    constraints = analysis.get("constraint_summary", {}) if isinstance(analysis, Mapping) else {}
    triggered = constraints.get("triggered_names", {}) if isinstance(constraints, Mapping) else {}
    flags: list[str] = []
    if isinstance(triggered, Mapping):
        for severity in ("hard_fail", "soft_warn", "evidence_only"):
            values = triggered.get(severity, [])
            if isinstance(values, list):
                flags.extend(str(v) for v in values)
    text = _interpretation_text(run)
    snr = _first_signal(analysis, "median_snr")
    peaks = None
    assignment = analysis.get("assignment_evidence", {}) if isinstance(analysis, Mapping) else {}
    if isinstance(assignment, Mapping) and assignment.get("assigned_peak_count") is not None:
        peaks = assignment.get("assigned_peak_count")
    if peaks is None:
        peaks = _parse_numeric_from_text(text, "nmr_peaks")
    frames = _parse_numeric_from_text(text, "frames")
    matrix_quality = _parse_numeric_from_text(text, "matrix_quality")
    cos_signal = _parse_numeric_from_text(text, "cos_signal")
    t_range = ""
    match = re.search(r"T_range=([^;]+)", text)
    if match:
        t_range = match.group(1)
    return {
        "run_id": str(run.get("run_id", "")),
        "manifest_ref": ref.replace("\\", "/"),
        "technique": technique,
        "mode": source_role or technique,
        "source_role": source_role,
        "source_refs": all_refs,
        "source_basenames": sorted(set(names)),
        "samples": samples,
        "package_evidence_ids": evidence_ids,
        "status": str(run.get("status", "unknown")),
        "figure_count": len(run.get("figures", []) if isinstance(run.get("figures", []), list) else []),
        "evidence_count": len(evidence_ids),
        "quality_flags": sorted(set(flags)),
        "median_snr": snr,
        "peak_count": peaks,
        "frames": int(frames) if frames is not None else None,
        "temperature_range": t_range,
        "matrix_quality": matrix_quality,
        "cos_signal": cos_signal,
        "interpretation": text,
        "raw_paths_hint": raw_paths,
    }


def _load_run_records(package_dir: Path) -> list[dict[str, Any]]:
    manifest = _json_load(package_dir / "manifest.json", {}) or {}
    refs = manifest.get("run_manifests", [])
    if not refs:
        refs = [f"runs/{p.name}" for p in sorted((package_dir / "runs").glob("*.json"))]
    records: list[dict[str, Any]] = []
    for ref in refs:
        path = package_dir / str(ref)
        run = _json_load(path)
        if not isinstance(run, Mapping):
            continue
        records.append(_run_record(run, str(ref), []))
    return sorted(records, key=lambda row: row["run_id"])


def _raw_match_kind(record: Mapping[str, Any], rel: str) -> str:
    """Return the provenance strength for one raw-file/run association."""

    rel_norm = _norm_path(rel)
    refs = [_norm_path(v) for v in record.get("source_refs", [])]
    for ref in refs:
        if ref.endswith("/" + rel_norm) or ref.endswith(rel_norm):
            return "explicit_path"
        # Directory artifact (notably the six IR sequences).
        if "/insu-ftir/" in ref and rel_norm.startswith("insu-ftir/"):
            if rel_norm.startswith(ref[ref.index("/insu-ftir/") + 1 :] + "/"):
                return "directory_context"
    return "none"


def _match_run_to_raw(record: Mapping[str, Any], rel: str) -> bool:
    """Match only explicit paths or declared directory artifacts.

    A basename-only fallback cross-links duplicated DSC exports in distinct
    directories, so it is deliberately excluded from auditable provenance.
    """

    return _raw_match_kind(record, rel) != "none"


def _raw_rows(project_root: Path, run_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path, _spec in _iter_raw_files(project_root):
        rel = path.relative_to(project_root).as_posix()
        info = classify_source(Path(rel))
        matched = [(r, _raw_match_kind(r, rel)) for r in run_records]
        package_runs = [r for r, kind in matched if kind != "none"]
        match_kinds = sorted({kind for _r, kind in matched if kind != "none"})
        evidence_ids = sorted({eid for r in package_runs for eid in r.get("package_evidence_ids", [])})
        statuses = sorted({str(r.get("status")) for r in package_runs})
        rows.append(
            {
                "relative_path": rel,
                "sample": info["sample"],
                "technique": info["technique"],
                "mode": info["mode"],
                "format": path.suffix.lower().lstrip("."),
                "size_bytes": path.stat().st_size,
                "sha256": _sha256(path),
                "export_family": info["export_family"],
                "use": info["use"],
                "exclusion_reason": info["exclusion_reason"],
                "package_run_ids": sorted({str(r["run_id"]) for r in package_runs}),
                "package_evidence_ids": evidence_ids,
                "review_status": ";".join(statuses) if statuses else "not_in_package",
                "package_membership": "explicit" if "explicit_path" in match_kinds else ("directory_context" if match_kinds else "no"),
                "package_match_method": ";".join(match_kinds) if match_kinds else "none",
            }
        )
    # Hash groups are explicit, so exact duplicate exports are not accidentally
    # treated as independent experimental replicates.
    by_hash: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        by_hash[row["sha256"]].append(row)
    for row in rows:
        group = by_hash[row["sha256"]]
        row["duplicate_hash_count"] = len(group)
        row["duplicate_hash_paths"] = [v["relative_path"] for v in group] if len(group) > 1 else []
    return sorted(rows, key=lambda row: row["relative_path"])


def _read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _derived_summaries(project_root: Path) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for key, rel, role in DERIVED_TABLES:
        path = project_root / rel
        rows = _read_csv(path)
        result[key] = {
            "path": rel,
            "role": role,
            "exists": path.exists(),
            "rows": len(rows),
            "columns": list(rows[0].keys()) if rows else [],
            "sha256": _sha256(path) if path.exists() else "",
            "approved_rows": sum(str(r.get("approved_for_manuscript", "")).lower() == "true" for r in rows),
            "data": rows,
        }
    return result


def _package_summary(package_dir: Path, run_records: list[dict[str, Any]]) -> dict[str, Any]:
    manifest = _json_load(package_dir / "manifest.json", {}) or {}
    technique_meta = manifest.get("techniques", {}) if isinstance(manifest, Mapping) else {}
    review = _json_load(package_dir / "review-decision.json", {}) or {}
    decisions = review.get("decisions", []) if isinstance(review, Mapping) else []
    pending = sum(str(v.get("decision", "")) == "pending" for v in decisions if isinstance(v, Mapping))
    return {
        "package_id": manifest.get("package_id", package_dir.name),
        "version": manifest.get("version"),
        "status": manifest.get("status", "unknown"),
        "package_hash": manifest.get("package_hash", ""),
        "run_count": len(manifest.get("run_ids", [])) if isinstance(manifest.get("run_ids", []), list) else len(run_records),
        "evidence_count": int(manifest.get("evidence_count", 0) or 0),
        "figure_count": int(manifest.get("figure_count", 0) or 0),
        "source_hash_count": len(manifest.get("source_hashes", [])) if isinstance(manifest.get("source_hashes", []), list) else 0,
        "techniques": {
            str(k): {
                "run_count": len(v.get("run_ids", [])) if isinstance(v, Mapping) else 0,
                "evidence_count": int(v.get("evidence_count", 0) or 0) if isinstance(v, Mapping) else 0,
                "statuses": list(v.get("statuses", [])) if isinstance(v, Mapping) else [],
            }
            for k, v in technique_meta.items()
        },
        "review_status": str(review.get("status", "pending")),
        "pending_review_decisions": pending,
        "limitations": list(manifest.get("limitations", [])),
    }


def _sample_technique_matrix(raw_rows: list[dict[str, Any]], run_records: list[dict[str, Any]], package: Mapping[str, Any]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for sample in SAMPLES:
        for technique in TECHNIQUES:
            if technique == "NMR-solution":
                raw = [
                    r for r in raw_rows
                    if r.get("sample") == sample
                    and r.get("technique") == "NMR"
                    and r.get("mode") == "solution_1d"
                ]
            elif technique == "NMR-solid":
                raw = [
                    r for r in raw_rows
                    if r.get("sample") == sample
                    and r.get("technique") == "NMR"
                    and r.get("mode") == "solid_1d"
                ]
            else:
                raw = [r for r in raw_rows if r.get("sample") == sample and r.get("technique") == technique]
            runs = [r for r in run_records if sample in r.get("samples", []) and str(r.get("technique")) == technique.lower().replace("-solution", "").replace("-solid", "")]
            # The package uses lower-case technique names and a single six-file
            # series run for SAXS/WAXS; normalize that representation here.
            if technique.startswith("NMR-"):
                expected = "nmr"
                runs = [r for r in run_records if sample in r.get("samples", []) and r.get("technique") == expected and ((technique == "NMR-solid") == (r.get("source_role") == "solid_1d" or any("solid-nmr" in _norm_path(v) for v in r.get("source_refs", []))))]
            else:
                expected = technique.lower()
                runs = [r for r in run_records if sample in r.get("samples", []) and str(r.get("technique", "")).lower() == expected]
            evidence = []
            for run in runs:
                evidence.extend(run.get("package_evidence_ids", []))
            if len(runs) == 1 and len(run_records) and technique in {"SAXS", "WAXS"}:
                # Each series evidence item is sample-specific even though the
                # source is one coherent run.
                evidence_count = 1
            else:
                evidence_count = len(set(evidence))
            statuses = sorted({str(r.get("status")) for r in runs})
            if raw and runs:
                coverage = "package_and_raw"
            elif raw:
                coverage = "raw_not_in_package"
            else:
                coverage = "missing"
            out.append(
                {
                    "sample": sample,
                    "technique": technique,
                    "raw_file_count": len(raw),
                    "package_run_count": len(runs),
                    "package_evidence_count": evidence_count,
                    "review_status": ";".join(statuses) if statuses else "not_in_package",
                    "coverage": coverage,
                    "use": "primary" if technique == "DSC" else ("supporting" if technique in {"IR", "WAXS", "SAXS"} else "audit_only"),
                }
            )
    return out


def build_data_audit(project_root: Path, package_dir: Path) -> dict[str, Any]:
    """Return the complete source/run/sample audit used by v002."""

    project_root = Path(project_root).expanduser().resolve()
    package_dir = Path(package_dir).expanduser().resolve()
    manifest = _json_load(package_dir / "manifest.json", {}) or {}
    run_records = _load_run_records(package_dir)
    raw_rows = _raw_rows(project_root, run_records)
    derived = _derived_summaries(project_root)
    package = _package_summary(package_dir, run_records)
    sample_matrix = _sample_technique_matrix(raw_rows, run_records, package)
    counts = Counter(row["technique"] for row in raw_rows)
    return {
        "project_root": str(project_root),
        "package_dir": str(package_dir),
        "package": package,
        "manifest_keys": sorted(manifest.keys()) if isinstance(manifest, Mapping) else [],
        "raw_file_count": len(raw_rows),
        "raw_counts_by_technique": dict(sorted(counts.items())),
        "source_rows": raw_rows,
        "run_rows": run_records,
        "sample_technique_rows": sample_matrix,
        "derived_tables": derived,
        "audit_notes": [
            "Exact hash duplicates are retained in the ledger and are not counted as independent experimental replicates.",
            "The -50 suffix is a user-declared nominal PTMG feed fraction; it is not treated as measured composition.",
            "Origin .opju files are retained as external context because the canonical reader marks the format unsupported.",
            "NMR-solid raw files cover the three neat PA samples only; no PA-50 solid-NMR source is present.",
            "The six IR directories contain CSV/SPC export families; the approved pairwise FTIR/2D-COS table is the writing source.",
        ],
    }


def _float(row: Mapping[str, Any], key: str, default: float = float("nan")) -> float:
    try:
        return float(row.get(key, default))
    except (TypeError, ValueError):
        return default


def _fmt(value: Any, digits: int = 3) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value) if value is not None else ""
    if number != number:  # NaN
        return "—"
    return f"{number:.{digits}f}".rstrip("0").rstrip(".")


def _range_for(rows: list[Mapping[str, Any]], key: str, digits: int = 3) -> str:
    values = [_float(row, key) for row in rows]
    values = [v for v in values if v == v]
    if not values:
        return "—"
    return f"{_fmt(min(values), digits)}–{_fmt(max(values), digits)}"


def _pairwise_range_summary(rows: list[Mapping[str, Any]]) -> dict[str, str]:
    """Return source-derived ranges for one neat-PA/PA-50 five-point pair."""

    return {
        "pure_t_half": _range_for(rows, "t_half_pure_min"),
        "pa50_t_half": _range_for(rows, "t_half_PA_50_min"),
        "time_ratio": _range_for(rows, "R_t_t_half_PA_50_over_pure"),
        "rate_ratio": _range_for(rows, "R_G_PA_50_over_pure"),
    }


def _run_ids_for(run_records: list[Mapping[str, Any]], technique: str, sample: str | None = None) -> list[str]:
    result = []
    for run in run_records:
        if str(run.get("technique", "")).lower() != technique.lower():
            continue
        if sample and sample not in run.get("samples", []):
            continue
        result.extend(str(v) for v in run.get("package_evidence_ids", []))
    return sorted(set(result))


def _paired_dsc_evidence_ids(run_records: list[Mapping[str, Any]], hard_segment: str) -> list[str]:
    """Return package evidence for both members of one DSC pair.

    The interpretation claims combine the non-isothermal thermal context with
    the isothermal rate comparison.  Therefore their provenance must include
    both the neat and ``-50`` runs (including duplicate, separately persisted
    export families) rather than only the neat control selected by
    ``_run_ids_for``.
    """

    values: set[str] = set()
    for sample in (hard_segment, f"{hard_segment}-50"):
        values.update(_run_ids_for(run_records, "dsc", sample))
    return sorted(values)


def build_claim_matrix(audit: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Build a claim ledger with explicit evidence level and boundaries."""

    runs = list(audit.get("run_rows", []))
    package = audit.get("package", {})
    def ev(tech: str, sample: str | None = None) -> str:
        values = _run_ids_for(runs, tech, sample)
        return ";".join(values) if values else "none"

    def paired_dsc_ev(hard_segment: str) -> str:
        values = _paired_dsc_evidence_ids(runs, hard_segment)
        return ";".join(values) if values else "none"

    rows = [
        ("C01", "observed", "PA6-50 在五个相对窗口位置的主结晶半时间均短于纯 PA6。", "analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv", ev("dsc", "PA6"), "R04;R07", "仅限当前样品批次、热史和 ΔTrel=0–4 °C 操作窗口；不等同于相同热力学过冷度。", "可写‘在本窗口内加快’；不可写成 PTMG 对所有 PA6 配方必然加快。"),
        ("C02", "observed", "PA12-50 在五个相对窗口位置的主结晶半时间均短于纯 PA12。", "analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv", ev("dsc", "PA12"), "R11;R12;R13", "绝对等温温度与纯 PA12 不同；表观速率比不是等过冷度倍数。", "可写‘当前窗口内加快’；不可归因于编号偶数本身。"),
        ("C03", "observed", "PA11-50 在五个相对窗口位置均呈较长的主结晶半时间。", "analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv", ev("dsc", "PA11"), "R08;R09;R10", "PA11 是当前配对反向对照，不代表全部奇数尼龙。", "可写‘当前配对中减慢’；不可写成奇数尼龙普遍减慢。"),
        ("C04", "compared", "PA6、PA12 的加快和 PA11 的减慢贯穿各自完整五点轨迹。", "analysis_output/control_dsc/dsc_qc_review.csv", ev("dsc"), "R02;R04;R20", "五点是温度轨迹，不是五个独立批次重复；未作显著性推断。", "不可写成统计显著或普适规律。"),
        ("C05", "observed", "ΔTrel=Tiso−Tc,min 只用于对齐每个样品自身有效窗口的位置。", "analysis_output/control_dsc/dsc_kinetics.csv; analysis_output/paper_data/dsc_kinetics.csv", ev("dsc"), "R26", "Tc,min 是操作定义，不是平衡熔点或共同热力学过冷度。", "不可写成 equal-supercooling 比较。"),
        ("C06", "observed", "PTMG 名义引入后六个样品的非等温 Tm、Tc 与质量归一化焓发生硬段依赖的变化。", "analysis_output/paper_figures_pairwise/nonisothermal_dsc_thermal_summary.csv", ev("dsc"), "R01;R04", "焓按总样品质量，不能直接换算硬段结晶度。", "不可混用旧积分表或宣称实测组成。"),
        ("C07", "interpreted", "PA6-50 在较小的非等温表观 Tm−Tc 差值下仍加快，单纯表观温差增大不足以解释该方向。", "analysis_output/paper_figures_pairwise/nonisothermal_dsc_thermal_summary.csv + analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv", paired_dsc_ev("PA6"), "R04;R07", "Tm 不是平衡熔点；迁移/松弛只是相容解释。", "不可写成已测得扩散系数提高。"),
        ("C08", "interpreted", "PA12-50 的较大表观温差与链段重排收益可能叠加，因而加速幅度最大。", "analysis_output/paper_figures_pairwise/nonisothermal_dsc_thermal_summary.csv + analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv", paired_dsc_ev("PA12"), "R11;R12;R13", "无法分离热学背景和迁移率贡献。", "不可报告贡献率或本征倍数。"),
        ("C09", "interpreted", "PA11-50 在较大表观温差下仍减慢，提示界面/规则堆砌代价可能压过松弛收益。", "analysis_output/paper_figures_pairwise/nonisothermal_dsc_thermal_summary.csv + analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv", paired_dsc_ev("PA11"), "R08;R10;R16", "尚未直接测量配准代价或限速步骤。", "不可写成唯一分子机制。"),
        ("C10", "observed", "统一 X=0.05–0.80 拟合得到的 Avrami n 记录主结晶区曲线形状。", "analysis_output/control_dsc/dsc_kinetics.csv; analysis_output/paper_data/dsc_kinetics.csv", ev("dsc"), "R26", "n 同时受成核分布、生长、相区限制和碰撞影响。", "不可把 n 等同唯一成核方式或生长维数。"),
        ("C11", "observed", "纯 PA 的 n 随窗口位置变化较大，单一指数不能解释全部速率方向。", "analysis_output/control_dsc/dsc_kinetics.csv", ev("dsc"), "R26", "仅描述表观拟合形状。", "不可用 n 排出唯一机理顺序。"),
        ("C12", "supporting", "专项成对降温 FTIR/2D-COS 中 N–H、酰胺 I/II 存在共同且不完全同步响应。", "analysis_output/paper_data/ftir_2dcos_pairwise_summary.csv", ev("ir"), "R05;R23;R24;R25", "专项表可用于支持性讨论；通用 IR provider 仍 paper_ready=False。", "不可仅凭二维符号指定唯一氢键或严格先后。"),
        ("C13", "supporting", "PA-50 的酰胺与 PTMG 醚键邻域在温度扰动下共同变化。", "ftir_2dcos_pairwise_nh_ether_PA*-50.npz", ev("ir"), "R05;R15", "相关性不等于特定 N–H···O 物种计数。", "不可报告氢键种类比例。"),
        ("C14", "audit", "六个通用 IR 温度序列已入包但因序列/赋值门槛保留为 review_required。", str(package.get("package_id", "")) + "/manifest.json", ev("ir"), "R05;R25", "band_support=0、trend_reproducible=False、paper_ready=False。", "不可把 provider 自动输出升级为新定量结果。"),
        ("C15", "supporting", "归一化 WAXS/SAXS 曲线显示硬段有序堆砌和中低 q 轮廓发生差异化变化。", "analysis_output/paper_data/scattering_profiles.csv", ev("waxs") + ";" + ev("saxs"), "R10;R27", "无独立背景扣除；绝对量、晶粒尺寸和可靠长周期不可用。", "不可报告 Xc、Scherrer 尺寸或绝对长周期。"),
        ("C16", "audit", "26 个 NMR runs 覆盖六个样品的溶液 1D 和三种纯 PA 的固体导出家族，但审核仍 pending。", str(package.get("package_id", "")) + "/review-decision.json", ev("nmr"), "R01;R06;R15", "固体 NMR 无 PA-50 文件；存在 assignment_limited、low-SNR 和零峰条目。", "不可由 NMR 单独证明组成、分子量、结晶/非晶比例或因果。"),
        ("C17", "interpreted", "迁移率/构象松弛、连续硬段规则堆砌和软硬段界面约束的竞争模型与联合证据相容。", "C01–C16 联合证据", ev("dsc") + ";" + ev("ir") + ";" + ev("nmr") + ";" + ev("waxs") + ";" + ev("saxs"), "R04;R07;R16;R20", "工作模型，不做参数反演，不声称唯一性。", "不可写成已揭示唯一分子机制。"),
        ("C18", "compared", "PA11 的反向响应排除了‘只要增加柔性就必然加快’的单因子解释。", "dsc_pure_vs_pa50_pairwise.csv + C12/C15", ev("dsc", "PA11"), "R02;R08;R20", "这是模型约束，不是奇偶定律证明。", "不可外推到全部奇数尼龙。"),
        ("A01", "unverified", "隐藏的真实软段含量、分子量或连续硬段长度差异可能参与速率方向。", "数据审计与缺失项清单", ev("nmr"), "—", "当前没有定量组成、GPC 或批次重复。", "必须在后续实验中核验，不可当作已知原因。"),
        ("A02", "unverified", "不同样品的热力学驱动力差异可能改变表观速率比。", "非等温 DSC + ΔTrel 定义", ev("dsc"), "R08;R12;R13", "缺少平衡熔点或统一热力学过冷度。", "不可称为本征速率倍数。"),
        ("A03", "unverified", "二维 IR、NMR 或散射处理伪影可能夸大结构差异。", "各技术 QC/审查状态", ev("ir") + ";" + ev("nmr") + ";" + ev("waxs") + ";" + ev("saxs"), "R23;R25;R27", "需人工审核、背景校正和原位重复。", "不可用低质量指标支持因果。"),
    ]
    result: list[dict[str, Any]] = []
    for claim_id, level, claim, source, evidence, refs, boundary, wording in rows:
        result.append(
            {
                "claim_id": claim_id,
                "evidence_level": level,
                "claim": claim,
                "source_paths": source,
                "package_evidence_ids": evidence,
                "literature_refs": refs,
                "boundary": boundary,
                "allowed_wording": wording,
                "sample_scope": ";".join(SAMPLES),
                "review_status": package.get("status", "unknown"),
            }
        )
    return result


def _literature_rows(project_root: Path) -> list[dict[str, str]]:
    source = project_root / "manuscript" / "even-nylon-crystallization-v001" / "literature_matrix.csv"
    rows = _read_csv(source)
    for row in rows:
        row["v002_use"] = row.get("evidence_role", "")
        row["citation_check"] = "pending"
        if row.get("ref_id") == "R20":
            row["year_basis"] = "print issue 2009; online-first 2008"
        elif row.get("ref_id") == "R13":
            row["year_basis"] = "matrix issue year 2021; online metadata also recorded"
        else:
            row["year_basis"] = "local matrix"
    return rows


def _markdown_table(headers: list[str], rows: Iterable[Iterable[Any]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    for row in rows:
        values = [str(v).replace("|", "\\|").replace("\n", " ") for v in row]
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def _pairwise_tables(audit: Mapping[str, Any]) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    derived = audit.get("derived_tables", {})
    pair = list(derived.get("isothermal_dsc_pairwise", {}).get("data", []))
    # The approved kinetics table is deliberately split by source family:
    # neat PA under control_dsc and PA-50 under paper_data.  Preserve both
    # provenance records in the audit while using their combined rows only for
    # all-six-sample descriptive summaries.
    kin = list(derived.get("isothermal_dsc_kinetics", {}).get("data", []))
    kin.extend(derived.get("isothermal_dsc_kinetics_pa50", {}).get("data", []))
    noniso = list(derived.get("nonisothermal_dsc_summary", {}).get("data", []))
    return pair, kin, noniso


def _nmr_rows(audit: Mapping[str, Any]) -> list[dict[str, Any]]:
    return [row for row in audit.get("run_rows", []) if str(row.get("technique", "")).lower() == "nmr"]


def _ir_rows(audit: Mapping[str, Any]) -> list[dict[str, Any]]:
    return [row for row in audit.get("run_rows", []) if str(row.get("technique", "")).lower() == "ir"]


def _scattering_rows(audit: Mapping[str, Any], technique: str) -> list[dict[str, Any]]:
    return [row for row in audit.get("run_rows", []) if str(row.get("technique", "")).lower() == technique.lower()]


def _approved_thermal_gap(row: Mapping[str, Any]) -> float:
    """Prefer the reviewed table gap; reconstruct it only when the field is absent."""

    approved = _float(row, "DeltaT_C")
    if approved == approved:
        return approved
    return _float(row, "Tm_C") - _float(row, "Tc_C")


def _dsc_source_role_counts(audit: Mapping[str, Any]) -> dict[str, int]:
    """Count package DSC runs by their persisted source role."""

    roles = {"isothermal_dsc": 0, "nonisothermal_dsc": 0}
    for row in audit.get("run_rows", []):
        if str(row.get("technique", "")).lower() != "dsc":
            continue
        role = str(row.get("source_role", ""))
        if role in roles:
            roles[role] += 1
    return roles


FIGURES: tuple[dict[str, str], ...] = (
    {
        "id": "F1",
        "title": "纯 PA 与 PA-50 的非等温 DSC 热曲线",
        "path": "analysis_output/paper_figures_pairwise/Fig2_nonisothermal_dsc_thermal_curves.png",
        "role": "main",
        "source": "Fig2_nonisothermal_dsc_thermal_curves.png + nonisothermal_dsc_thermal_summary.csv",
        "boundary": "偏移热流曲线只作热学背景；焓不换算硬段结晶度",
    },
    {
        "id": "F2",
        "title": "纯 PA 与 PA-50 的等温 DSC 候选放热曲线",
        "path": "analysis_output/paper_figures_pairwise/Fig3_pure_pa_vs_pa50_exotherms.png",
        "role": "main",
        "source": "analysis_output/paper_figures_pairwise/Fig3_pure_pa_vs_pa50_exotherms.png + analysis_output/control_dsc/dsc_candidate_curves.csv + analysis_output/paper_data/dsc_candidate_curves.csv",
        "boundary": "候选放热曲线来自等温 candidate_curves；用于显示轨迹，不替代配对 t1/2 速率表",
    },
    {
        "id": "F3",
        "title": "纯 PA 与 PA-50 的等温相对结晶度",
        "path": "analysis_output/paper_figures_pairwise/Fig4_pure_pa_vs_pa50_relative_crystallinity.png",
        "role": "main",
        "source": "Fig4_pure_pa_vs_pa50_relative_crystallinity.png + two approved dsc_kinetics.csv tables",
        "boundary": "相对结晶度用于主结晶区形状比较，不用于绝对结晶度",
    },
    {
        "id": "F4",
        "title": "纯 PA 与 PA-50 的等温主结晶动力学",
        "path": "analysis_output/paper_figures_pairwise/Fig8_pure_pa_vs_pa50_kinetics.png",
        "role": "main",
        "source": "Fig8_pure_pa_vs_pa50_kinetics.png + dsc_pure_vs_pa50_pairwise.csv + two dsc_kinetics.csv tables",
        "boundary": "配对窗口位置不是等热力学过冷度；t1/2 与 G 才是速率指标",
    },
    {
        "id": "F5",
        "title": "成对 FTIR 局部谱带响应",
        "path": "analysis_output/paper_figures_pairwise/Fig6_pairwise_ftir_peak_response.png",
        "role": "main",
        "source": "Fig6_pairwise_ftir_peak_response.png + 专项 FTIR/2D-COS 输出",
        "boundary": "样品内趋势；不报告氢键比例",
    },
    {
        "id": "F6",
        "title": "WAXS/SAXS 归一化结构轮廓",
        "path": "analysis_output/paper_figures_pairwise/Fig9_pure_pa_vs_pa50_scattering.png",
        "role": "main",
        "source": "外部派生附件：analysis_output/paper_figures_pairwise/Fig9_pure_pa_vs_pa50_scattering.png + analysis_output/paper_data/scattering_profiles.csv",
        "boundary": "图和曲线是证据包之外的外部派生附件；无独立背景扣除，仅作归一化定性显示，不报告绝对量或定量形貌",
    },
    {
        "id": "F7",
        "title": "成对 FTIR/2D-COS 相关域",
        "path": "analysis_output/paper_figures_pairwise/Fig7_pairwise_ftir_2dcos_nh_amide_i.png",
        "role": "main",
        "source": "Fig7_pairwise_ftir_2dcos_nh_amide_i.png + ftir_2dcos_pairwise_summary.csv + NPZ",
        "boundary": "不依据符号指定唯一氢键或严格顺序",
    },
    {
        "id": "S2",
        "title": "NMR 溶液/固体谱（12 溶液；14 固体，仅审计；review-only）",
        "path": "analysis_output/paper_figures_pairwise/Fig1_readable_offset_solution_nmr.png",
        "role": "supplement",
        "source": "all-data package: 12 solution NMR runs + 14 solid NMR runs",
        "boundary": "12 个溶液 NMR 可供来源/化学环境复核；14 个固体 NMR 仅审计（仅纯 PA、无 PA-50 对照）；人工归属前不承担组成/机制结论",
    },
)


def _figure_marker(figure: Mapping[str, Any]) -> str:
    return f"<!--FIGURE:{figure['id']}|{figure['path']}|{figure['title']}|{figure['role']}-->"


def _render_manuscript_markdown(audit: Mapping[str, Any], claims: list[dict[str, Any]], literature: list[dict[str, str]]) -> str:
    pair, kin, noniso = _pairwise_tables(audit)
    package = audit.get("package", {})
    matrix = audit.get("sample_technique_rows", [])
    raw_rows = list(audit.get("source_rows", []))
    dsc_runs = [row for row in audit.get("run_rows", []) if str(row.get("technique", "")).lower() == "dsc"]
    dsc_run_roles = Counter(str(row.get("source_role", "")).lower() for row in dsc_runs)
    dsc_isothermal_runs = dsc_run_roles.get("isothermal_dsc", 0)
    dsc_nonisothermal_runs = dsc_run_roles.get("nonisothermal_dsc", 0)
    dsc_isothermal_raw = sum(row.get("mode") == "isothermal_dsc" for row in raw_rows)
    dsc_nonisothermal_raw = sum(row.get("mode") == "nonisothermal_dsc" for row in raw_rows)
    dsc_isothermal_evidence = len(
        {item for row in dsc_runs if str(row.get("source_role", "")).lower() == "isothermal_dsc" for item in row.get("package_evidence_ids", [])}
    )
    dsc_nonisothermal_evidence = len(
        {item for row in dsc_runs if str(row.get("source_role", "")).lower() == "nonisothermal_dsc" for item in row.get("package_evidence_ids", [])}
    )
    ir_rows = _ir_rows(audit)
    nmr_rows = _nmr_rows(audit)
    pair_by = {s: [r for r in pair if r.get("hard_segment") == s] for s in ("PA6", "PA11", "PA12")}
    pair_stats = {sample: _pairwise_range_summary(rows) for sample, rows in pair_by.items()}
    kin_by = {s: [r for r in kin if r.get("sample") == s] for s in SAMPLES}
    n_ranges = {s: _range_for(kin_by[s], "n") for s in SAMPLES}
    thermal_gaps = {
        str(row.get("sample", "")): _approved_thermal_gap(row)
        for row in noniso
    }
    dsc_roles = _dsc_source_role_counts(audit)
    dsc_run_count = sum(dsc_roles.values())

    lines: list[str] = []
    lines.extend(
        [
            "# PTMG 引入后偶数编号尼龙基热塑性弹性体结晶加快的动力学来源：全数据证据约束的 PA6、PA12 与 PA11 对照",
            "",
            "**English title:** Kinetic origin of accelerated crystallization in even-numbered-nylon-based thermoplastic elastomers after PTMG incorporation: an all-data comparison of PA6, PA12, and PA11",
            "",
            f"**稿件状态：** 中文工作稿 v002（基于 `{package.get('package_id', 'unknown')}` v{package.get('version', '')}；包状态 `{package.get('status', 'unknown')}`；作者、单位和目标期刊待补）。",
            "",
            "## 摘要",
            "",
            f"聚四亚甲基醚二醇（PTMG）同时改变聚酰胺硬段附近的构象松弛、连续可结晶序列和软硬段界面。本文利用同一项目中的全量可追溯数据，比较纯 PA6、PA11、PA12 与对应 PTMG 名义投料 50 wt% 的 PA6-50、PA11-50、PA12-50。等温 DSC 在每个样品自身具有完整主放热峰的五点窗口内，以 ΔTrel=Tiso−Tc,min 对齐窗口位置，并以 t1/2、G=1/t1/2 和 X=0.05–0.80 区间的表观 Avrami 拟合描述主结晶过程。PA6-50 与纯 PA6 的 t1/2 比值为 {pair_stats['PA6']['time_ratio']}，PA12-50 与纯 PA12 为 {pair_stats['PA12']['time_ratio']}；PA11-50 则为 {pair_stats['PA11']['time_ratio']}，方向相反。非等温 DSC 提供峰位和质量归一化焓的热学背景；专项变温 FTIR/二维相关红外、溶液/固体 NMR、WAXS 与 SAXS 被逐项审计并按证据强度分层。FTIR/2D-COS 支持酰胺和醚键邻域共同重排，NMR 记录化学环境与文件质量但尚待人工归属，归一化散射曲线只支持堆砌和相区轮廓变化。综合结果支持一个工作模型：PTMG 的迁移/松弛收益与连续硬段缩短、规则堆砌受扰和界面重排代价竞争；在当前样品和热史下，松弛收益可能在 PA6 中占优，PA12 还叠加较大的表观热学温差，而 PA11 的界面—堆砌代价更突出。该结论限于三组当前配对体系，不构成普适偶奇定律。",
            "",
            "**关键词：** 热塑性聚酰胺弹性体；PTMG；等温结晶；半结晶时间；链段松弛；界面约束；二维相关红外；NMR；WAXS；SAXS",
            "",
            "## Abstract",
            "",
            f"PTMG incorporation can increase local segmental relaxation while shortening continuous crystallizable polyamide sequences and increasing soft–hard interfaces. We compare neat PA6, PA11, and PA12 with paired PA6-50, PA11-50, and PA12-50 samples, where -50 denotes a nominal PTMG feed fraction of 50 wt% rather than a measured composition. Isothermal DSC trajectories were aligned at ΔTrel=Tiso−Tc,min within each sample’s own five-point valid window. PA6-50 and PA12-50 accelerated across all paired positions, with t1/2 ratios of {pair_stats['PA6']['time_ratio']} and {pair_stats['PA12']['time_ratio']}, respectively, whereas PA11-50 slowed with ratios of {pair_stats['PA11']['time_ratio']}. Nonisothermal DSC, specialized cooling FTIR/two-dimensional correlation spectra, solution/solid NMR inventories, and qualitative normalized WAXS/SAXS profiles were used as bounded thermal, local-structure, and morphology context. The complete audit preserves {package.get('run_count', '')} package runs and {package.get('evidence_count', '')} evidence items, while review-required and unsupported quantities remain non-causal. The observations are consistent with competition among mobility gain, hard-segment packing disruption, and interfacial reorganization. This is a working model for the present paired samples, not a universal odd–even law.",
            "",
            "**Keywords:** thermoplastic polyamide elastomer; PTMG; isothermal crystallization; half-crystallization time; segmental relaxation; two-dimensional correlation spectroscopy; NMR; WAXS; SAXS",
            "",
            "## 1 引言",
            "",
            "热塑性聚酰胺弹性体（TPAE）由可结晶聚酰胺硬段和柔性聚醚软段共同构成。硬段晶区及其强缔合区提供可逆物理交联，软段和软硬段界面则控制链段输运、相区重排和加工热史[1,3–6,17–22]。因此，PTMG 的引入并不是单一的‘增柔’操作：它可能降低非晶链段的松弛阻力，也可能打断连续酰胺序列、改变氢键配准和晶体堆砌。已有 PA6、PA11 和 PA12 研究分别显示，过冷度、热史、氢键和晶型选择会共同影响结晶动力学[2,4,7–16]。",
            "",
            "本研究的配对数据呈现出一个非单调结果：PA6-50 和 PA12-50 在各自有效等温窗口中均加快，而 PA11-50 在同一配对规则下减慢。这个反向对照不支持‘柔性提高就必然加快’的单因子解释，但三种硬段仍不足以把编号奇偶性从重复单元长度、酰胺密度、分子量和热史中分离出来。本文的研究问题因此限定为：在当前名义组成与测试程序下，哪些相互竞争的热学、局部结构和界面因素与 PA6/PA12 的加快及 PA11 的减慢相容？",
            "",
            "为避免不同入口产生不同科学结论，本文以冻结的 all-data evidence package 为来源边界。等温 DSC 是唯一直接的速率证据；非等温 DSC 是热学背景；专项 FTIR/2D-COS、NMR、WAXS 和 SAXS 逐项进入审计、结果讨论或补充材料。所有自动诊断量、未审核峰归属和无背景绝对散射量均保持其原有 review-only 或 diagnostic-only 状态。",
            "",
            "## 2 样品、数据来源与处理方法",
            "",
            "### 2.1 样品命名与证据包",
            "",
            "研究对象为 PA6、PA6-50、PA11、PA11-50、PA12 和 PA12-50。‘-50’仅表示用户声明的 PTMG 名义投料质量分数，不表示通过定量 NMR、元素分析或质量平衡测得的实际软段含量。冻结包 `" + str(package.get("package_id", "")) + "` v" + str(package.get("version", "")) + "` 的状态为 `" + str(package.get("status", "")) + "`，包含 " + str(package.get("run_count", "")) + " 个 runs、" + str(package.get("evidence_count", "")) + " 个 evidence items、" + str(package.get("figure_count", "")) + " 个逻辑图和 " + str(package.get("source_hash_count", "")) + " 个来源哈希；package hash 为 `" + str(package.get("package_hash", "")) + "`。所有人工审核决定仍为 pending。",
            "",
            "### 2.2 全量数据审计规则",
            "",
            f"原始目录共审计 {audit.get('raw_file_count', 0)} 个文件：DSC 等温/非等温文本与 Origin 项目、六个 IR 温度目录、12 个溶液 NMR CSV、14 个固体 NMR CSV、6 个 WAXS RAW 和 6 个 SAXS EDF。每个文件记录相对路径、大小、SHA-256、样品标签、技术、导出家族、包内 run 反查结果和用途。相同 SHA-256 的导出保留在台账中，但不作为独立重复。不能由规范化文件名推断实测组成。",
            "",
            _markdown_table(
                ["技术/模式", "原始文件数", "包内 runs", "包内 evidence", "写作角色"],
                [
                    ["DSC 等温文本", dsc_isothermal_raw, dsc_isothermal_runs, dsc_isothermal_evidence, "速率主证据（表格）"],
                    ["DSC 非等温文本/Origin", dsc_nonisothermal_raw, dsc_nonisothermal_runs, dsc_nonisothermal_evidence, "热学背景；.opju 仅审计"],
                    ["IR 温度序列", sum(r.get("raw_file_count", 0) for r in matrix if r["technique"] == "IR"), package.get("techniques", {}).get("ir", {}).get("run_count", 0), package.get("techniques", {}).get("ir", {}).get("evidence_count", 0), "专项摘要支持；通用 provider 审计"],
                    ["NMR 溶液/固体", sum(r.get("raw_file_count", 0) for r in matrix if r["technique"].startswith("NMR")), package.get("techniques", {}).get("nmr", {}).get("run_count", 0), package.get("techniques", {}).get("nmr", {}).get("evidence_count", 0), "结构文件审计/补充材料"],
                    ["WAXS", sum(r.get("raw_file_count", 0) for r in matrix if r["technique"] == "WAXS"), package.get("techniques", {}).get("waxs", {}).get("run_count", 0), package.get("techniques", {}).get("waxs", {}).get("evidence_count", 0), "归一化定性"],
                    ["SAXS", sum(r.get("raw_file_count", 0) for r in matrix if r["technique"] == "SAXS"), package.get("techniques", {}).get("saxs", {}).get("run_count", 0), package.get("techniques", {}).get("saxs", {}).get("evidence_count", 0), "归一化定性"],
                ],
            ),
            "",
            "六个样品的覆盖矩阵和逐文件清单分别见 `data_audit.md` 与 `data_audit.csv`。其中，固体 NMR 只有纯 PA6/PA11/PA12，没有 PA-50 固体谱；六个样品均有溶液 NMR。DSC 的等温文本存在多个目录/导出家族，非等温 Origin `.opju` 因格式不支持保留为外部上下文，不伪装为 canonical run。",
            "",
            "### 2.3 DSC：热学背景与等温动力学",
            "",
            "非等温参数采用指定汇总表 `analysis_output/paper_figures_pairwise/nonisothermal_dsc_thermal_summary.csv`，程序为 10 °C min⁻¹ 的升温—降温—再升温。等温主结晶分别采用 `analysis_output/control_dsc/dsc_kinetics.csv`（纯 PA，15 行）和 `analysis_output/paper_data/dsc_kinetics.csv`（PA-50，15 行），并以 `dsc_pure_vs_pa50_pairwise.csv` 进行 15 个位置的配对比较。每个样品的 Tc,min 定义为恒温段仍可获得完整主放热峰的最低有效温度；ΔTrel=Tiso−Tc,min 只对齐样品内窗口位置。",
            "",
            "相对结晶度由放热峰累计积分得到，t1/2 为 X=0.5 的时间，G=1/t1/2。表观 Avrami 方程为 X=1−exp(−ktⁿ)，所有平台固定在 X=0.05–0.80 区间拟合。由于 k 的量纲随 n 改变，跨样品的方向只用 t1/2 和 G 判断；n 仅描述主结晶区形状。",
            "",
            "### 2.4 FTIR/二维相关红外、NMR 与散射",
            "",
            "专项 FTIR/2D-COS 采用降温序列，统计表中的谱数为 PA6/PA6-50=23/21、PA11/PA11-50=20/18、PA12/PA12-50=19/16。N–H、酰胺 I/II 和 PTMG 醚键窗口用于判断共同变化和不完全同步响应；没有根据单一二维符号指定唯一氢键物种或严格运动先后。通用 IR provider 的六个温度序列虽然已转换入包，但 sequence_axis_incomplete、cross_peak_without_band_assignment 和 noda_rule_not_applicable 等门槛使其保持 review_required。",
            "",
            "NMR 包含 12 个溶液 1D runs（六个样品的 ¹H/¹³C 文件族）和 14 个固体 1D runs（三个纯 PA 的 ¹H/¹³C 文本/重复导出）。存在 assignment_limited、nmr_low_snr、nmr_low_peak_count 和零峰条目；所以 NMR 用于溯源、文件质量和作者后续人工核对，不用于定量组成、晶区/非晶区比例或因果机制。",
            "",
            "WAXS/SAXS 只采用 `background_method=none; curves normalized for display only` 的轮廓。WAXS 的峰数和非单调序列警告、SAXS 的低 q 空洞/污染和 PA12 q* 风险均列入审计；不报告绝对强度、结晶度、晶粒尺寸或可靠长周期。",
            "",
            "## 3 结果与讨论",
            "",
            "### 3.1 全量覆盖与质量边界",
            "",
            f"全量审计的目的不是把所有自动指标都提升为科学结果，而是让每个来源都有明确去向。DSC 的 {dsc_run_count} 个包内 runs（{dsc_roles['isothermal_dsc']} 个等温来源和 {dsc_roles['nonisothermal_dsc']} 个非等温文本来源）支撑主要热分析；六个 IR 序列、26 个 NMR runs、6 个 WAXS 和 6 个 SAXS evidence items 均保留包内来源和审核状态。所有包内 evidence items 当前为 review_required，故本文对低层级数据使用‘支持、相容、提示’等措辞。",
            "",
            _markdown_table(
                ["样品", "DSC", "IR", "NMR-溶液", "NMR-固体", "WAXS", "SAXS"],
                [
                    [s] + [next((r for r in matrix if r["sample"] == s and r["technique"] == t), {}).get("coverage", "missing") for t in ("DSC", "IR", "NMR-solution", "NMR-solid", "WAXS", "SAXS")]
                    for s in SAMPLES
                ],
            ),
            "",
            "### 3.2 非等温 DSC：PTMG 改变了热学环境，但没有决定速率方向",
            "",
            _markdown_table(
                ["样品", "Tm / °C", "Tc / °C", "Tm−Tc / °C", "ΔHm / J g⁻¹", "ΔHc / J g⁻¹"],
                [
                    [r.get("sample", ""), _fmt(r.get("Tm_C")), _fmt(r.get("Tc_C")), _fmt(_approved_thermal_gap(r)), _fmt(r.get("DeltaHm_J_g")), _fmt(r.get("DeltaHc_J_g"))]
                    for r in noniso
                ],
            ),
            "",
            f"六组样品的 Tm 和 Tc 均因 PTMG 名义引入而下移，但幅度依硬段而异。PA6 的 Tm−Tc 由 {_fmt(thermal_gaps.get('PA6'))} °C 变为 {_fmt(thermal_gaps.get('PA6-50'))} °C，PA11 由 {_fmt(thermal_gaps.get('PA11'))} °C 变为 {_fmt(thermal_gaps.get('PA11-50'))} °C，PA12 由 {_fmt(thermal_gaps.get('PA12'))} °C 变为 {_fmt(thermal_gaps.get('PA12-50'))} °C。PA11-50 与 PA12-50 都有较大的表观差值，却分别减慢和加快，说明该差值只能提供热学背景，不能单独解释净速率。质量归一化焓同时受软段稀释和可结晶硬段比例影响，不被解释为绝对结晶度。",
            "",
            _figure_marker(FIGURES[0]),
            "",
            "### 3.3 等温 DSC：PA6、PA12 加快，PA11 反向减慢",
            "",
            "表 2 列出全部 15 个成对位置，而不是只报告区间。",
            "",
            _figure_marker(FIGURES[1]),
            "",
            _figure_marker(FIGURES[2]),
            "",
            _markdown_table(
                ["硬段", "ΔTrel / °C", "Tiso 纯 / °C", "Tiso PA-50 / °C", "t1/2 纯 / min", "t1/2 PA-50 / min", "时间比", "G 比", "方向"],
                [
                    [r.get("hard_segment", ""), r.get("delta_trel_C", ""), r.get("tiso_pure_C", ""), r.get("tiso_PA_50_C", ""), _fmt(r.get("t_half_pure_min")), _fmt(r.get("t_half_PA_50_min")), _fmt(r.get("R_t_t_half_PA_50_over_pure")), _fmt(r.get("R_G_PA_50_over_pure")), r.get("kinetic_change", "")]
                    for r in pair
                ],
            ),
            "",
            f"PA6-50 的 t1/2 为 {pair_stats['PA6']['pa50_t_half']} min，纯 PA6 为 {pair_stats['PA6']['pure_t_half']} min，五点时间比为 {pair_stats['PA6']['time_ratio']}，G 比为 {pair_stats['PA6']['rate_ratio']}。PA12-50 的 t1/2 为 {pair_stats['PA12']['pa50_t_half']} min，纯 PA12 为 {pair_stats['PA12']['pure_t_half']} min，时间比为 {pair_stats['PA12']['time_ratio']}，G 比为 {pair_stats['PA12']['rate_ratio']}。PA11-50 的 t1/2 为 {pair_stats['PA11']['pa50_t_half']} min，纯 PA11 为 {pair_stats['PA11']['pure_t_half']} min，时间比为 {pair_stats['PA11']['time_ratio']}，G 比为 {pair_stats['PA11']['rate_ratio']}。",
            "",
            "这些比值回答的是‘在各自有效窗口相同相对位置，轨迹向前还是向后移动’，不是‘相同热力学过冷度下的本征速率倍数’。五个位置属于一条温度轨迹，不能替代独立批次重复或统计显著性。",
            "",
            _figure_marker(FIGURES[3]),
            "",
            "### 3.4 表观 Avrami 指数：记录形状，不承担唯一机理",
            "",
            _markdown_table(
                ["样品", "n 范围", "用途"],
                [[s, n_ranges[s], "主结晶区曲线形状"] for s in SAMPLES],
            ),
            "",
            f"PA6-50、PA11-50 和 PA12-50 的 n 分别为 {n_ranges['PA6-50']}、{n_ranges['PA11-50']} 和 {n_ranges['PA12-50']}；纯 PA6、PA11、PA12 分别为 {n_ranges['PA6']}、{n_ranges['PA11']} 和 {n_ranges['PA12']}。PA6-50 与 PA11-50 的 n 区间接近而速率方向相反，说明单一 n 不能解释净速率。对于软硬段共存的体系，n 还会受成核时间分布、相区限制、晶体碰撞和可结晶硬段供给影响。",
            "",
            "### 3.5 FTIR/二维相关红外：局部酰胺—醚键邻域共同重排",
            "",
            _markdown_table(
                ["样品", "包内序列 frames", "温度范围 / °C", "matrix quality", "cos signal", "provider 状态"],
                [[r.get("samples", [""])[0] if r.get("samples") else "", r.get("frames", "—"), r.get("temperature_range", "—"), _fmt(r.get("matrix_quality")), _fmt(r.get("cos_signal")), "review_required; paper_ready=False"] for r in ir_rows],
            ),
            "",
            "专项成对降温 FTIR/2D-COS 中 N–H、酰胺 I 和酰胺 II 具有同步与异步相关域；PA-50 还可观察到酰胺与 PTMG 醚键窗口共同变化。该结果支持‘局部缔合—构象调整—界面邻域重排耦合’的结构图景，但相关峰不是氢键物种的直接计数。由于通用 provider 的 band support、序列轴和 Noda 规则门槛尚未通过，正文不采用其自动峰数作为新的定量结论。",
            "",
            _figure_marker(FIGURES[4]),
            "",
            _figure_marker(FIGURES[6]),
            "",
            "### 3.6 NMR：完整入包，限定为结构文件与审核证据",
            "",
            _markdown_table(
                ["run", "样品/文件", "模式", "median SNR", "峰数", "质量/归属标记", "用途"],
                [[r.get("run_id", ""), ", ".join(r.get("source_basenames", [])[:2]), "solution/solid 1D", _fmt(r.get("median_snr")), _fmt(r.get("peak_count"), 0), ", ".join(r.get("quality_flags", [])[:4]) or "—", "review-only"] for r in nmr_rows],
            ),
            "",
            "NMR 审计的完整表在 `data_audit.csv` 和 `supplementary_data_index.md` 中。高 SNR 文件可供作者核对化学环境，低 SNR、零峰和 assignment_limited 文件则保留为待处理项。尤其不能把固体 NMR 的通用区域积分直接写成晶区/非晶区比例，也不能以溶液 NMR 自动峰数证明 PTMG 实际含量、分子量或端基计量。",
            "",
            _figure_marker(FIGURES[7]),
            "",
            "### 3.7 WAXS/SAXS：有序堆砌和相区组织的定性背景",
            "",
            _markdown_table(
                ["技术", "样品覆盖", "包内 evidence", "主要 QC 限制", "正文角色"],
                [
                    ["WAXS", "PA6/PA6-50/PA11/PA11-50/PA12/PA12-50", "6", "peak_count_insufficient; amorphous_partition_unstable; size_without_multi_peak_support", "定性峰形/峰位"],
                    ["SAXS", "PA6/PA6-50/PA11/PA11-50/PA12/PA12-50", "6", "low_q_void_dominant；PA12 q* 低 q 污染", "定性中低 q 轮廓"],
                ],
            ),
            "",
            "PTMG 引入后 WAXS 峰形以及 SAXS 中低 q 轮廓发生差异化变化，与连续硬段和软硬段界面同时重排相容。由于曲线没有独立背景扣除且仅作显示归一化，峰强不能换算绝对结晶度，2π/q* 也不作为可靠长周期。静态散射只能说明最终结构组织不同，不能从单个轮廓反演主结晶过程中的限速步骤。",
            "",
            _figure_marker(FIGURES[5]),
            "",
            "### 3.8 机制解释：迁移率—堆砌—界面竞争",
            "",
            "当前数据最适合用一个竞争模型组织，而不是用编号奇偶性直接作因果变量。PTMG 可能缩短非晶/界面链段的松弛时间，使硬段更容易采样可结晶构象；同时，它也缩短连续聚酰胺序列、增加连接点附近的不规则性，并改变酰胺缔合和晶格配准。软硬段界面既可能提供局部输运或异相成核环境，也可能限制取向和重排。因此，净结晶时间是迁移/松弛收益、规则堆砌惩罚和界面贡献的耦合结果。",
            "",
            f"对 PA6，PA6-50 在较小的非等温表观 Tm−Tc 差值下仍在五点轨迹加快，说明单纯‘更大表观温差’不足以解释方向；在工作模型中，界面构象松弛收益可能占优。对 PA12，PA12-50 的非等温表观 Tm−Tc 差值由 {_fmt(thermal_gaps.get('PA12'))} °C 变为 {_fmt(thermal_gaps.get('PA12-50'))} °C，且获得最大 G 比 {pair_stats['PA12']['rate_ratio']}；这与热学背景和迁移收益叠加的模型相容，但无法拆分贡献。对 PA11，PA11-50 虽有较大表观差值却在全部五点减慢，提示连续硬段规则堆砌和界面重排代价可能更突出。FTIR/2D-COS 与散射轮廓的共同变化为这一解释提供结构旁证，NMR 则暂不承担因果选择。",
            "",
            "两个替代解释必须保留：其一，真实软段含量、分子量和连续硬段长度可能并未被当前名义配方控制；其二，不同样品的平衡熔点和热力学驱动力可能不同。另有测量/处理伪影的可能性，特别是未审核的通用 IR、固体 NMR 峰归属和未背景校正的散射。因此，本文把竞争模型标记为‘与数据相容的工作假设’，而不是已证实的唯一机制。",
            "",
            "## 4 局限性与可检验预测",
            "",
            "1. **热力学比较：** ΔTrel 不是共同过冷度。应测定平衡熔点或采用 Hoffman–Weeks/快速扫描方法，在统一驱动力下重做配对动力学。若 PA6/PA12 的方向仍保持，迁移/界面解释更有说服力；若幅度收敛，热学背景贡献更大。",
            "2. **组成与分子量：** 补充定量 NMR、元素分析、GPC/SEC、硬段序列或扩链状态表征，并进行独立合成批次重复。当前 NMR 数据尚不能替代这些测量。",
            "3. **动力学分解：** 用原位 WAXS/SAXS、偏光显微和成核密度/生长速率分析区分成核与晶体生长；用流变或介电谱检验 PA6-50/PA12-50 是否在结晶温区具有更有利的松弛时间尺度。",
            "4. **谱学与散射审核：** 完成人工 NMR 峰归属、IR 序列轴和 Noda 规则复核，并对 SAXS/WAXS 做背景、几何和绝对强度校准；在此之前，所有相关指标继续保留 review-only。",
            "5. **统计设计：** 当前五点是温度轨迹而非独立重复。投稿前需增加批次/样品重复、误差条和预先定义的比较模型。",
            "",
            "## 5 结论",
            "",
            f"1. 在各样品自身完整五点等温窗口的相对位置上，PA6-50 和 PA12-50 的主结晶轨迹整体前移，PA11-50 整体后移；时间比范围分别为 {pair_stats['PA6']['time_ratio']}、{pair_stats['PA12']['time_ratio']} 和 {pair_stats['PA11']['time_ratio']}。",
            "2. 非等温 DSC 显示 PTMG 改变了六组样品的热学环境，但 PA11 与 PA12 的方向相反，表明表观峰位差不能单独决定速率。",
            "3. 专项 FTIR/2D-COS 与归一化 WAXS/SAXS 对局部谱带响应、堆砌与相区轮廓提供限定性结构旁证；NMR 仅作为来源、文件质量和人工归属待审计的证据，不单独承担结构重排或因果论证。",
            "4. 最符合现有证据的工作模型是迁移/松弛收益与规则堆砌、界面重排代价的竞争。该模型解释 PA11 反例，但尚未隔离组成、分子量、热力学驱动力和具体限速步骤，不能外推为普适偶奇定律。",
            "",
            "## 数据可用性、审核状态与补充材料",
            "",
            f"原始数据、分析表和证据包均保留在项目目录 `{audit.get('project_root', '')}`。本稿随附 `data_audit.csv`（逐文件 SHA-256 与包内反查）、`claim_evidence_literature_matrix.csv`（{len(claims)} 条主张/替代解释）、`figure_table_supplement_index.md` 和 `supplementary_data_index.md`。all-data package 的人工审核为 `{package.get('review_status', 'pending')}`，待审核决定 {package.get('pending_review_decisions', '—')} 条；文献 DOI 与书目信息沿用本地矩阵，在线核验状态仍为 `citation_check=pending`。",
            "",
            "## 参考文献",
            "",
        ]
    )
    for index, row in enumerate(literature, 1):
        citation = row.get("citation", "").strip()
        doi = row.get("doi", "").strip()
        lines.append(f"[{index}] {citation} DOI: {doi}.（{row.get('verification_status', 'citation_check=pending')}；v002 citation_check=pending）")
    return "\n".join(lines) + "\n"


def _write_csv(path: Path, rows: list[Mapping[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            converted = dict(row)
            for key, value in list(converted.items()):
                if isinstance(value, list):
                    converted[key] = ";".join(str(v) for v in value)
                elif isinstance(value, dict):
                    converted[key] = json.dumps(value, ensure_ascii=False, sort_keys=True)
            writer.writerow(converted)


def _write_data_audit_files(out: Path, audit: Mapping[str, Any]) -> None:
    source_fields = [
        "relative_path", "sample", "technique", "mode", "format", "size_bytes", "sha256",
        "export_family", "duplicate_hash_count", "duplicate_hash_paths", "use", "exclusion_reason",
        "package_membership", "package_match_method", "package_run_ids", "package_evidence_ids", "review_status",
    ]
    _write_csv(out / "data_audit.csv", list(audit.get("source_rows", [])), source_fields)
    run_fields = [
        "run_id", "manifest_ref", "technique", "mode", "source_role", "source_refs", "source_basenames", "samples",
        "package_evidence_ids", "status", "evidence_count", "figure_count", "quality_flags", "median_snr",
        "peak_count", "frames", "temperature_range", "matrix_quality", "cos_signal", "interpretation",
    ]
    _write_csv(out / "run_audit.csv", list(audit.get("run_rows", [])), run_fields)
    matrix_fields = ["sample", "technique", "raw_file_count", "package_run_count", "package_evidence_count", "review_status", "coverage", "use"]
    _write_csv(out / "sample_technique_coverage.csv", list(audit.get("sample_technique_rows", [])), matrix_fields)

    package = audit.get("package", {})
    lines = [
        "# 全量数据审计（v002）",
        "",
        f"- 项目根目录：`{audit.get('project_root', '')}`",
        f"- 证据包：`{audit.get('package_dir', '')}`",
        f"- package_id/version：`{package.get('package_id', '')}` / `v{package.get('version', '')}`",
        f"- package 状态：`{package.get('status', '')}`；hash：`{package.get('package_hash', '')}`",
        f"- 包内 runs/evidence/figures/source hashes：{package.get('run_count', '')}/{package.get('evidence_count', '')}/{package.get('figure_count', '')}/{package.get('source_hash_count', '')}",
        f"- 人工审核：`{package.get('review_status', '')}`；pending={package.get('pending_review_decisions', '')}",
        f"- 原始文件总数：{audit.get('raw_file_count', 0)}",
        "",
        "## 样品 × 技术覆盖",
        "",
        _markdown_table(
            ["样品", "技术", "原始文件", "包内 runs", "包内 evidence", "覆盖", "审核状态", "写作角色"],
            [[r.get("sample", ""), r.get("technique", ""), r.get("raw_file_count", ""), r.get("package_run_count", ""), r.get("package_evidence_count", ""), r.get("coverage", ""), r.get("review_status", ""), r.get("use", "")] for r in audit.get("sample_technique_rows", [])],
        ),
        "",
        "## 派生表",
        "",
        _markdown_table(
            ["键", "路径", "角色", "行数", "approved 行", "SHA-256"],
            [[k, v.get("path", ""), v.get("role", ""), v.get("rows", ""), v.get("approved_rows", ""), v.get("sha256", "")] for k, v in audit.get("derived_tables", {}).items()],
        ),
        "",
        "## 审计解释",
    ]
    lines.extend(f"- {note}" for note in audit.get("audit_notes", []))
    lines.extend(
        [
            "",
            "逐文件来源、哈希、导出家族和包内 run 反查见 `data_audit.csv`；逐 run 质量/审核字段见 `run_audit.csv`；六个样品的覆盖矩阵见 `sample_technique_coverage.csv`。",
            "",
            "### 需要人工处理的来源",
            "",
            "- 所有 62 个 evidence items 的 review decision 当前为 pending。",
            "- 通用 IR 温度序列的 sequence axis、关键谱带支持和 Noda 规则仍未通过。",
            "- NMR 中部分文件峰数为零或信噪比低，且固体 NMR 没有 PA-50 对照。",
            "- WAXS/SAXS 的背景、绝对校准和峰支持不足以给出定量形貌参数。",
            "- `.opju` 只作为外部上下文保留，不当作已 canonical 化的分析 run。",
        ]
    )
    (out / "data_audit.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_claim_matrix(out: Path, claims: list[dict[str, Any]]) -> None:
    fields = ["claim_id", "evidence_level", "claim", "source_paths", "package_evidence_ids", "literature_refs", "boundary", "allowed_wording", "sample_scope", "review_status"]
    _write_csv(out / "claim_evidence_literature_matrix.csv", claims, fields)


def _write_literature_matrix(out: Path, literature: list[dict[str, str]]) -> None:
    fields = list(literature[0].keys()) if literature else ["ref_id", "citation", "doi", "year", "verification_status"]
    _write_csv(out / "literature_matrix.csv", literature, fields)


def _write_figure_index(out: Path) -> None:
    lines = [
        "# 图表与补充材料索引（v002）",
        "",
        "正文图按结论优先排序；低质量或待审核来源不被隐藏，而是以补充材料/审计角色保留。",
        "",
        "## 正文候选图",
        "",
        _markdown_table(["图号", "文件", "内容/来源", "角色", "边界"], [[f["id"], f["path"], f["source"], f["role"], f["boundary"]] for f in FIGURES if f["role"] == "main"]),
        "",
        "## 补充图",
        "",
        _markdown_table(["图号", "文件", "内容/来源", "角色", "边界"], [[f["id"], f["path"], f["source"], f["role"], f["boundary"]] for f in FIGURES if f["role"] != "main"]),
        "",
        "## 表格",
        "",
        _markdown_table(
            ["表号", "路径/文件", "内容", "准入"],
            [
                ["表 1", "analysis_output/paper_figures_pairwise/nonisothermal_dsc_thermal_summary.csv", "六样品 Tm/Tc/焓", "指定汇总"],
                ["表 2", "analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv", "全部 15 个配对位置与速率比", "主定量"],
                ["表 S1", "analysis_output/control_dsc/dsc_kinetics.csv + analysis_output/paper_data/dsc_kinetics.csv", "30 个平台的 n/k/R²/拟合窗口（纯 PA 与 PA-50 各 15 个）", "主结晶区形状"],
                ["表 S2", "analysis_output/paper_data/ftir_2dcos_pairwise_summary.csv", "谱数、温度范围和匹配规则", "专项支持"],
                ["表 S3", "data_audit.csv + run_audit.csv", "554 个原始文件与 52 个 runs", "全量审计"],
                ["表 S4", "sample_technique_coverage.csv", "六样品 × 六技术覆盖", "全量审计"],
            ],
        ),
        "",
        "## 使用规则",
        "",
        "1. 图 2–4 的定量值来自当前指定 CSV；不混用 `origin_native_dsc` 旧积分事件。",
        "2. 图 5、7 的二维红外只报告共同/不完全同步响应，不由符号推导唯一氢键或严格先后。",
        "3. 图 6 的 WAXS/SAXS 仅作归一化显示；背景方法和绝对量限制必须写入图题。",
        "4. 图 S2 及包内 NMR/通用 IR provider SVG 均显式标为 `review-only`，直至人工 review decision 更新。",
    ]
    (out / "figure_table_supplement_index.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_supplement_index(out: Path, audit: Mapping[str, Any]) -> None:
    nmr = _nmr_rows(audit)
    ir = _ir_rows(audit)
    lines = [
        "# 补充材料数据索引（v002）",
        "",
        "本索引把‘全部数据’分成可直接复核的表、支持性图和仅审计来源；不把 review-only 指标升级为主张。",
        "",
        "## 补充表",
        "",
        "- **S1 等温 DSC 全量平台：** `analysis_output/control_dsc/dsc_kinetics.csv`（纯 PA，15 行）与 `analysis_output/paper_data/dsc_kinetics.csv`（PA-50，15 行）；合计 30 个平台，每个平台含 t1/2、G、n、k、R²、拟合区间和质量标志。",
        "- **S2 等温 DSC 质量审查：** `analysis_output/control_dsc/dsc_qc_review.csv`（平台点数、温度标准差、积分边界和基线方法）。",
        "- **S3 FTIR/2D-COS：** `analysis_output/paper_data/ftir_2dcos_pairwise_summary.csv`（六样品及专项输出文件名）。",
        "- **S4 散射摘要：** `analysis_output/paper_data/scattering_profile_summary.csv`（仅定性索引；不报告长周期为可靠定量）。",
        "- **S5 原始来源审计：** `data_audit.csv`（逐文件哈希、重复家族和包内反查）。",
        "- **S6 运行审计：** `run_audit.csv`（52 个 package runs 的来源、审核状态和触发限制）。",
        "",
        "## IR 温度序列",
        "",
        _markdown_table(["样品", "frames", "温度范围", "matrix quality", "cos signal", "状态"], [[r.get("samples", [""])[0] if r.get("samples") else "", r.get("frames", "—"), r.get("temperature_range", "—"), _fmt(r.get("matrix_quality")), _fmt(r.get("cos_signal")), "review-only"] for r in ir]),
        "",
        "## NMR 逐 run 审计",
        "",
        _markdown_table(["run", "文件", "SNR", "峰数", "限制", "状态"], [[r.get("run_id", ""), ", ".join(r.get("source_basenames", [])[:3]), _fmt(r.get("median_snr")), _fmt(r.get("peak_count"), 0), ", ".join(r.get("quality_flags", [])[:5]) or "—", "review-only"] for r in nmr]),
        "",
        "固体 NMR 文件族中的 `TXT`、`-70` 和同样本重复导出保留为导出家族，不自动视为独立重复；PA-50 固体 NMR 缺失是覆盖缺口，不用纯样谱替代。",
        "",
        "## 待补实验/审核",
        "",
        "- 完成 all-data package 的 62 项人工审核决定；",
        "- 完成 NMR 峰归属和零峰/低 SNR 文件复核；",
        "- 统一热力学过冷度并增加独立批次重复；",
        "- 对散射做背景、几何和绝对强度校准；",
        "- 对通用 IR provider 补齐序列轴、关键谱带和 Noda 规则证据。",
    ]
    (out / "supplementary_data_index.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _citation_dicts(literature: list[dict[str, str]]) -> list[dict[str, Any]]:
    return [
        {
            "version": 1,
            "citation_id": f"citation-{row.get('ref_id', index):s}",
            "key": row.get("ref_id", f"R{index:02d}"),
            "locator": row.get("doi", "") or f"local-matrix:{row.get('ref_id', index):s}",
            "mode": "static",
            "metadata": {"citation": row.get("citation", ""), "year": row.get("year", ""), "verification_status": "citation_check=pending"},
        }
        for index, row in enumerate(literature, 1)
    ]


def _claim_contract_dicts(claims: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for claim in claims:
        evidence = tuple(v for v in str(claim.get("package_evidence_ids", "")).split(";") if v and v != "none")
        result.append(
            {
                "version": 1,
                "claim_id": claim["claim_id"],
                "text": claim["claim"],
                "evidence_ids": list(evidence) or [claim.get("source_paths", "audit")],
                "metric_ids": [f"source:{claim.get('source_paths', '')}"],
                "figure_ids": [],
                "role": "hypothesis" if claim.get("evidence_level") in {"interpreted", "unverified"} else ("discussion" if claim.get("evidence_level") in {"supporting", "audit", "compared"} else "results"),
                "citation_keys": [v for v in str(claim.get("literature_refs", "")).split(";") if v and v != "—"],
                "evidence_level": claim.get("evidence_level", ""),
                "boundary": claim.get("boundary", ""),
            }
        )
    return result


def _figure_contract_dicts(project_root: Path) -> list[dict[str, Any]]:
    result = []
    for fig in FIGURES:
        path = project_root / fig["path"]
        result.append(
            {
                "version": 1,
                "figure_id": fig["id"],
                "title": fig["title"],
                "layout": "single",
                "source_ids": [fig["source"]],
                "metric_ids": [],
                "outputs": ["svg", "json", "png", "source"],
                "status": "candidate" if path.exists() else "missing_source_asset",
                "source_path": fig["path"],
                "role": fig["role"],
                "boundary": fig["boundary"],
            }
        )
    return result


def _manuscript_json(audit: Mapping[str, Any], claims: list[dict[str, Any]], literature: list[dict[str, str]], markdown: str, docx_status: str) -> dict[str, Any]:
    from polynexus.suite.paper_contracts import ManuscriptSource

    package = audit.get("package", {})
    source_id = "manuscript-" + hashlib.sha256(f"{package.get('package_id')}:{package.get('version')}:v002".encode()).hexdigest()[:16]
    claim_contracts = _claim_contract_dicts(claims)
    figures = _figure_contract_dicts(Path(audit.get("project_root", ".")))
    citations = _citation_dicts(literature)
    formulas = [
        {"version": 1, "formula_id": "formula-t-half", "expression": "G = 1 / t_1/2", "variables": {"G": "reciprocal rate index", "t_1/2": "time at X=0.5"}, "units": {"G": "min^-1", "t_1/2": "min"}, "source": "analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv"},
        {"version": 1, "formula_id": "formula-avrami", "expression": "X = 1 - exp(-k t^n)", "variables": {"X": "relative crystallinity", "k": "apparent rate parameter", "n": "apparent shape exponent", "t": "time"}, "units": {"X": "1", "k": "fit-dependent", "n": "1"}, "source": "Avrami 1939; dsc_kinetics.csv"},
        {"version": 1, "formula_id": "formula-delta-trel", "expression": "DeltaTrel = Tiso - Tc,min", "variables": {"Tiso": "isothermal setpoint", "Tc,min": "lowest valid isothermal temperature"}, "units": {"DeltaTrel": "C"}, "source": "dsc_kinetics.csv and operating definition"},
    ]
    limitations = [
        *[str(value) for value in package.get("limitations", []) if str(value).strip()],
        "five-point trajectories are not independent replicates",
        "solid NMR lacks PA-50 pair",
        "literature citation_check=pending",
    ]
    sections = [
        {"name": "摘要", "claim_ids": ["C01", "C02", "C03", "C17"]},
        {"name": "方法与数据审计", "claim_ids": ["C05", "C06", "C14", "C16"]},
        {"name": "结果与讨论", "claim_ids": [f"C{i:02d}" for i in range(1, 19)]},
        {"name": "局限性与可检验预测", "claim_ids": ["A01", "A02", "A03"]},
    ]
    contract = ManuscriptSource.create(
        package_id=str(package.get("package_id", "unknown")) or "unknown",
        claim_ids=tuple(str(value["claim_id"]) for value in claim_contracts),
        figure_plan_ids=tuple(str(value["figure_id"]) for value in figures),
        citation_ids=tuple(str(value["citation_id"]) for value in citations),
        limitations=tuple(limitations),
    ).to_dict()
    working_draft = {
        "title_zh": "PTMG 引入后偶数编号尼龙基热塑性弹性体结晶加快的动力学来源：全数据证据约束的 PA6、PA12 与 PA11 对照",
        "title_en": "Kinetic origin of accelerated crystallization in even-numbered-nylon-based thermoplastic elastomers after PTMG incorporation: an all-data comparison of PA6, PA12, and PA11",
        "research_question": "在各样品自身有效等温结晶窗口中，PTMG 引入为何使 PA6 和 PA12 的主结晶加快，而使 PA11 减慢？",
        "language": "zh-CN",
        "paper_type": "evidence-grounded original research article",
        "package": package,
        "scope": {
            "samples": list(SAMPLES),
            "nominal_ptmg_feed_fraction": "50 wt% for -50 samples; measured fraction not established",
            "comparison_protocol": "paired pure-PA/PA-50 at DeltaTrel=0-4 C within each sample's own valid window",
            "primary_technique": "isothermal DSC",
            "supporting_techniques": ["nonisothermal DSC", "specialized FTIR/2D-COS", "solution/solid NMR audit", "qualitative WAXS", "qualitative SAXS"],
            "non_goals": ["universal odd-even law", "equal thermodynamic supercooling claim", "exact composition/molecular weight", "unique hydrogen-bond species", "absolute scattering quantities"],
        },
        "sections": sections,
        "data_audit_summary": {
            "raw_file_count": audit.get("raw_file_count", 0),
            "raw_counts_by_technique": audit.get("raw_counts_by_technique", {}),
            "run_count": package.get("run_count", 0),
            "evidence_count": package.get("evidence_count", 0),
            "claim_count": len(claims),
            "literature_count": len(literature),
        },
        "markdown_sha256": hashlib.sha256(markdown.encode("utf-8")).hexdigest(),
        "docx_status": docx_status,
        "author_review_required": True,
        "created_at": str(date.today()),
    }
    contract.update({
        "schema": "polynexus.manuscript-working-draft.v2",
        "draft_version": "v002",
        "status": "working_draft",
        "source_id": source_id,
        "language": working_draft["language"],
        "paper_type": working_draft["paper_type"],
        "title_zh": working_draft["title_zh"],
        "title_en": working_draft["title_en"],
        "research_question": working_draft["research_question"],
        "package": package,
        "scope": working_draft["scope"],
        "sections": sections,
        "claims": claim_contracts,
        "figures": figures,
        "citations": citations,
        "formulas": formulas,
        "formula_ids": [str(value["formula_id"]) for value in formulas],
        "data_audit_summary": working_draft["data_audit_summary"],
        "markdown_sha256": working_draft["markdown_sha256"],
        "docx_status": working_draft["docx_status"],
        "author_review_required": working_draft["author_review_required"],
        "created_at": working_draft["created_at"],
        "projection": {"working_draft": working_draft},
    }
    )
    return contract


def _strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = text.replace("\\|", "|")
    return text


def _set_run_font(run: Any, name: str = "SimSun", size: int = 10) -> None:
    from docx.shared import Pt
    from docx.oxml.ns import qn

    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)


def _set_style_font(style: Any, name: str = "SimSun", size: float = 10.5) -> None:
    """Set both Western and East-Asian font faces on a base Word style."""

    from docx.oxml.ns import qn
    from docx.shared import Pt

    style.font.name = name
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)


def _add_docx_table(document: Any, rows: list[list[str]]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0 and tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = _strip_md(str(value))
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, "SimSun", 8 if len(rows[0]) > 6 else 9)
            if i == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    document.add_paragraph()


def _build_docx(markdown: str, output: Path, project_root: Path) -> str:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches
    except Exception as exc:  # pragma: no cover - dependency optional
        return f"unavailable:{exc.__class__.__name__}"

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    for style_name in ("Normal", "Body Text"):
        try:
            _set_style_font(styles[style_name])
        except Exception:
            pass

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        marker = re.match(r"<!--FIGURE:([^|]+)\|([^|]+)\|([^|]+)\|([^>]+)-->", line.strip())
        if marker:
            fig_id, rel_path, caption, role = marker.groups()
            image_path = project_root / rel_path
            if image_path.exists():
                try:
                    document.add_picture(str(image_path), width=Inches(6.15))
                    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    document.add_paragraph(f"[图 {fig_id} 图片插入失败：{rel_path}]")
            else:
                document.add_paragraph(f"[图 {fig_id}：{caption}；源文件未找到：{rel_path}]")
            p = document.add_paragraph(f"图 {fig_id}  {caption}（{role}; source={rel_path}）")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                _set_run_font(run, "SimSun", 9)
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                if not re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$", lines[index]):
                    values = [v.strip() for v in lines[index].strip().strip("|").split("|")]
                    table_lines.append(values)
                index += 1
            _add_docx_table(document, table_lines)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            p = document.add_heading(_strip_md(heading.group(2)), level=min(level, 3))
            for run in p.runs:
                _set_run_font(run, "SimSun", 15 if level == 1 else (13 if level == 2 else 11))
            index += 1
            continue
        if line.startswith("- ") or re.match(r"^\d+\.\s", line):
            text = re.sub(r"^(-|\d+\.)\s+", "", line)
            p = document.add_paragraph(_strip_md(text), style="List Bullet" if line.startswith("-") else "List Number")
        else:
            p = document.add_paragraph(_strip_md(line))
        for run in p.runs:
            _set_run_font(run)
        index += 1
    document.save(output)
    return "created"


def _has_prohibited_causal_overclaim(markdown: str) -> bool:
    """Detect an unqualified causal claim without rejecting its limitation text.

    The manuscript must explicitly mention prohibited extrapolations in order to
    rule them out.  A raw keyword search therefore incorrectly marks sentences
    such as ``不构成普适偶奇定律`` as failures.  Examine the sentence containing
    each protected phrase and only fail when it lacks a clear limiting cue.
    """

    protected = re.compile(r"唯一分子机制|普适偶奇定律|相同热力学过冷度下的本征")
    limit_cue = re.compile(
        r"不构成|不能|不可|不是|并非|尚未|未证明|未证实|未能|未被|"
        r"无(?:法|从)|不等同|不外推|不代表|不支持|不作为|不应|无法"
    )
    for match in protected.finditer(markdown):
        before = markdown[: match.start()]
        after = markdown[match.end() :]
        sentence_start = max(before.rfind("。"), before.rfind("！"), before.rfind("？"), before.rfind("\n")) + 1
        endings = [index for index in (after.find("。"), after.find("！"), after.find("？"), after.find("\n")) if index >= 0]
        sentence_end = match.end() + (min(endings) if endings else len(after))
        if not limit_cue.search(markdown[sentence_start:sentence_end]):
            return True
    return False


def _custom_preflight(manuscript: Mapping[str, Any], markdown: str, audit: Mapping[str, Any], docx_status: str) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    try:
        from polynexus.suite.preflight import preflight_manuscript

        shared = preflight_manuscript(manuscript).to_dict()
        errors.extend(shared.get("errors", []))
        warnings.extend(shared.get("warnings", []))
    except Exception as exc:
        warnings.append(f"shared_preflight_unavailable:{exc.__class__.__name__}")
        shared = {}
    package = audit.get("package", {})
    derived = audit.get("derived_tables", {})
    checks: dict[str, Any] = {
        "shared_contract_preflight": shared,
        "package_status": {"status": "review_required" if package.get("status") == "review_required" else "review", "value": package.get("status")},
        "package_counts": {"runs": package.get("run_count"), "evidence": package.get("evidence_count"), "figures": package.get("figure_count"), "source_hashes": package.get("source_hash_count")},
        "raw_audit": {"status": "pass" if audit.get("raw_file_count", 0) else "review_required", "raw_file_count": audit.get("raw_file_count", 0)},
        "dsc_pairwise": {"status": "pass" if derived.get("isothermal_dsc_pairwise", {}).get("rows") == 15 else "review_required", "rows": derived.get("isothermal_dsc_pairwise", {}).get("rows", 0)},
        "dsc_kinetics": {
            "status": "pass"
            if (
                derived.get("isothermal_dsc_kinetics", {}).get("rows") == 15
                and derived.get("isothermal_dsc_kinetics_pa50", {}).get("rows") == 15
            )
            else "review_required",
            "rows": (
                derived.get("isothermal_dsc_kinetics", {}).get("rows", 0)
                + derived.get("isothermal_dsc_kinetics_pa50", {}).get("rows", 0)
            ),
        },
        "nonisothermal_summary": {"status": "pass" if derived.get("nonisothermal_dsc_summary", {}).get("rows") == 6 else "review_required", "rows": derived.get("nonisothermal_dsc_summary", {}).get("rows", 0)},
        "all_six_samples": {"status": "pass" if all(any(r.get("sample") == s for r in audit.get("sample_technique_rows", [])) for s in SAMPLES) else "review_required"},
        "causal_boundary": {"status": "pass" if not _has_prohibited_causal_overclaim(markdown) else "failed"},
        "docx": {"status": "pass" if docx_status == "created" else "review_required", "value": docx_status},
        "citation_check": {"status": "pending", "reference_count": len(manuscript.get("citations", [])), "note": "local metadata retained; online verification not claimed"},
    }
    if checks["causal_boundary"]["status"] == "failed":
        errors.append("causal_boundary_phrase_detected")
    if package.get("status") != "review_required":
        warnings.append("package_status_not_review_required")
    warnings.extend(["all evidence-package review decisions remain pending", "citation_check=pending", "scientific conclusions require human review"])
    status = "failed" if errors else "pass_with_human_review"
    return {
        "schema": "polynexus.manuscript-preflight.v2",
        "version": "v002",
        "run_date": str(date.today()),
        "manuscript": "manuscript.md",
        "overall_status": status,
        "errors": sorted(set(errors)),
        "warnings": sorted(set(warnings)),
        "checks": checks,
        "required_human_actions_before_submission": [
            "Complete all package review decisions, especially IR and NMR.",
            "Verify composition, molecular-weight distribution and independent batch replicates.",
            "Repeat kinetics under a common thermodynamic supercooling definition.",
            "Add calibrated/background-corrected scattering if quantitative morphology is claimed.",
            "Independently verify all literature metadata and DOI records.",
        ],
    }


def build_manuscript(project_root: Path, package_dir: Path, output_dir: Path) -> dict[str, str]:
    """Generate a new v002 directory; refuse to overwrite an existing one."""

    project_root = Path(project_root).expanduser().resolve()
    package_dir = Path(package_dir).expanduser().resolve()
    output_dir = Path(output_dir).expanduser().resolve()
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f"output directory is not empty: {output_dir}")
    if package_dir == output_dir or package_dir in output_dir.parents:
        raise ValueError("output directory cannot be inside immutable evidence package")
    output_dir.mkdir(parents=True, exist_ok=True)

    audit = build_data_audit(project_root, package_dir)
    package = audit.get("package", {})
    claims = build_claim_matrix(audit)
    literature = _literature_rows(project_root)
    markdown = _render_manuscript_markdown(audit, claims, literature)
    (output_dir / "manuscript.md").write_text(markdown, encoding="utf-8")
    _write_data_audit_files(output_dir, audit)
    _write_claim_matrix(output_dir, claims)
    _write_literature_matrix(output_dir, literature)
    _write_figure_index(output_dir)
    _write_supplement_index(output_dir, audit)

    brief = {
        "version": 1,
        "brief_id": "all-data-paper-v002",
        "research_question": "在各样品自身有效等温结晶窗口中，PTMG 引入为何使 PA6 和 PA12 的主结晶加快，而使 PA11 减慢？",
        "title_hint": "PTMG 引入后偶数编号尼龙基热塑性弹性体结晶加快的动力学来源",
        "comparison_scope": {"samples": list(SAMPLES), "paired_control": "neat PA versus corresponding PA-50", "window_alignment": "DeltaTrel=Tiso-Tc,min at 0-4 C", "primary_technique": "isothermal DSC", "supporting_techniques": ["nonisothermal DSC", "specialized FTIR/2D-COS", "solution/solid NMR audit", "qualitative WAXS", "qualitative SAXS"], "package": package_dir.name},
        "figure_budget": {"main_max": 7, "supporting_max": 8},
        "notes": "Use all sources through main evidence, bounded supporting evidence, supplementary material, or explicit audit/exclusion rows. Do not infer composition, molecular weight, unique hydrogen-bond species, absolute scattering quantities, or a universal odd-even law.",
        "needs_input": [],
        "status": "draft",
    }
    (output_dir / "paper-brief.json").write_text(json.dumps(brief, ensure_ascii=False, indent=2), encoding="utf-8")
    docx_path = output_dir / "manuscript.docx"
    docx_status = _build_docx(markdown, docx_path, project_root)
    manuscript = _manuscript_json(audit, claims, literature, markdown, docx_status)
    (output_dir / "manuscript.json").write_text(json.dumps(manuscript, ensure_ascii=False, indent=2), encoding="utf-8")
    preflight = _custom_preflight(manuscript, markdown, audit, docx_status)
    (output_dir / "preflight.json").write_text(json.dumps(preflight, ensure_ascii=False, indent=2), encoding="utf-8")
    review_lines = [
        "# v002 审核状态",
        "",
        f"- package：`{package.get('package_id', '')}` v{package.get('version', '')}; status=`{package.get('status', '')}`",
        f"- 包内 runs/evidence/figures：{package.get('run_count', '')}/{package.get('evidence_count', '')}/{package.get('figure_count', '')}",
        f"- pending review decisions：{package.get('pending_review_decisions', '')}",
        f"- 文献：{len(literature)} 条；`citation_check=pending`",
        f"- DOCX：`{docx_status}`",
        "",
        "## 投稿前必须完成",
        "",
        "1. 完成全部 evidence item 的人工科学审核，特别是通用 IR 序列轴/Noda 规则和 NMR 峰归属。",
        "2. 核验真实软段含量、分子量分布、连续硬段长度和独立批次重复。",
        "3. 在共同热力学过冷度下重复等温动力学，并区分成核和生长。",
        "4. 若报告散射定量形貌，补做背景/几何/绝对强度校准。",
        "5. 独立核验全部 27 条 DOI、年份、卷页和作者信息。",
        "",
        "v002 不修改原始数据、冻结证据包、v001/v002 旧稿或既有 DOCX。",
    ]
    (output_dir / "review-status.md").write_text("\n".join(review_lines) + "\n", encoding="utf-8")
    # Descriptive aliases make the external folder easy to find while the
    # canonical names above remain stable for CLI/ARS consumers.
    (output_dir / "TPAE_even_nylon_crystallization_mechanism_v002.md").write_text(markdown, encoding="utf-8")
    (output_dir / "TPAE_even_nylon_crystallization_mechanism_v002.json").write_text(json.dumps(manuscript, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"output_dir": str(output_dir), "markdown": str(output_dir / "manuscript.md"), "json": str(output_dir / "manuscript.json"), "docx": str(docx_path), "preflight": str(output_dir / "preflight.json"), "audit": str(output_dir / "data_audit.csv")}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", required=True, type=Path)
    parser.add_argument("--package", dest="package_dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args(argv)
    try:
        result = build_manuscript(args.project_root, args.package_dir, args.output)
    except (OSError, ValueError, FileExistsError) as exc:
        print(json.dumps({"status": "blocked", "error": str(exc)}, ensure_ascii=False, indent=2))
        return 2
    print(json.dumps({"status": "created", **result}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
