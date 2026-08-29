"""Evidence-first manuscript source assembly."""

from __future__ import annotations

from typing import Iterable, Any
import json
from pathlib import Path

from .paper_contracts import CitationRequest, ClaimRecord, FigurePlan, FormulaRecord, ManuscriptSource


def assemble_manuscript(
    source: ManuscriptSource,
    claims: Iterable[ClaimRecord] = (),
    figures: Iterable[FigurePlan] = (),
    citations: Iterable[CitationRequest] = (),
    formulas: Iterable[FormulaRecord] = (),
) -> dict[str, Any]:
    """Create a writing input projection; never changes scientific values."""
    claim_values = tuple(claims)
    figure_values = tuple(figures)
    citation_values = tuple(citations)
    formula_values = tuple(formulas)
    return {
        "version": 1, "source_id": source.source_id, "package_id": source.package_id,
        "sections": [
            {"name": "Results", "claim_ids": [c.claim_id for c in claim_values if c.role == "results"]},
            {"name": "Discussion", "claim_ids": [c.claim_id for c in claim_values if c.role != "results"]},
        ],
        "claims": [c.to_dict() for c in claim_values],
        "figures": [f.to_dict() for f in figure_values],
        "citations": [c.to_dict() for c in citation_values],
        "formulas": [f.to_dict() for f in formula_values],
        "limitations": list(source.limitations),
        "evidence_projection": dict(source.projection),
    }


def export_manuscript(manuscript: dict[str, Any], output_dir: str | Path) -> dict[str, str]:
    """Write non-destructive, reviewable manuscript artifacts."""
    root = Path(output_dir).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    json_path = root / "manuscript.json"
    md_path = root / "manuscript.md"
    json_path.write_text(json.dumps(manuscript, ensure_ascii=False, sort_keys=True, indent=2), encoding="utf-8")
    lines = ["# Manuscript draft", "", f"Package: `{manuscript.get('package_id', '')}`", ""]
    for section in manuscript.get("sections", []):
        lines.extend([f"## {section.get('name', 'Section')}", ""])
        for claim_id in section.get("claim_ids", []):
            claim = next((c for c in manuscript.get("claims", []) if c.get("claim_id") == claim_id), None)
            if claim:
                lines.extend([claim.get("text", ""), ""])
    md_path.write_text("\n".join(lines), encoding="utf-8")
    result = {"json": str(json_path), "markdown": str(md_path)}
    pdf_path = root / "manuscript.pdf"
    _write_minimal_pdf(pdf_path, lines)
    result["pdf"] = str(pdf_path)
    try:
        from docx import Document
        document = Document()
        document.add_heading("Manuscript draft", level=1)
        for section in manuscript.get("sections", []):
            document.add_heading(str(section.get("name", "Section")), level=2)
            for claim_id in section.get("claim_ids", []):
                claim = next((c for c in manuscript.get("claims", []) if c.get("claim_id") == claim_id), None)
                if claim:
                    document.add_paragraph(str(claim.get("text", "")))
        docx_path = root / "manuscript.docx"
        document.save(docx_path)
        result["docx"] = str(docx_path)
    except Exception:
        pass
    return result


def _write_minimal_pdf(path: Path, lines: list[str]) -> None:
    """Write a dependency-free, text-only PDF for layout smoke review."""
    visible = [line[:110] for line in lines[:55]] or [""]
    stream_lines = ["BT", "/F1 10 Tf", "50 780 Td"]
    for index, line in enumerate(visible):
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        if index:
            stream_lines.append("0 -14 Td")
        stream_lines.append(f"({escaped}) Tj")
    stream_lines.append("ET")
    stream = "\n".join(stream_lines).encode("latin-1", errors="replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode())
        payload.extend(obj)
        payload.extend(b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode())
    payload.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(payload)


__all__ = ["assemble_manuscript", "export_manuscript"]
