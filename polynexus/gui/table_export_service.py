"""Helpers for writing GUI table exports."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class TableExportData:
    headers: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]


def write_table_export(
    path,
    headers: list[str],
    matrix: list[list[str]],
    *,
    selected_filter: str = "",
) -> str:
    if not path or not headers or not matrix:
        return ""

    output_path = str(path)
    delimiter = "\t"
    lower_path = output_path.lower()
    if lower_path.endswith(".csv") or "csv" in str(selected_filter).lower():
        delimiter = ","
        if not lower_path.endswith(".csv"):
            output_path = f"{output_path}.csv"
    elif not lower_path.endswith(".tsv"):
        output_path = f"{output_path}.tsv"

    with open(output_path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh, delimiter=delimiter)
        writer.writerow(headers)
        writer.writerows(matrix)

    return output_path


def write_table_export_bundle(
    path,
    *,
    summary: TableExportData,
    full: TableExportData,
    diagnostics: TableExportData | None,
    selected_filter: str = "",
) -> list[str]:
    if (
        not path
        or not isinstance(summary, TableExportData)
        or not isinstance(full, TableExportData)
        or not summary.headers
        or not full.headers
    ):
        return []

    output_path = Path(path)
    lower_filter = str(selected_filter).lower()
    if output_path.suffix.lower() == ".xlsx" or any(
        marker in lower_filter for marker in ("excel", "xlsx")
    ):
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")
        _write_xlsx_bundle(output_path, summary, full, diagnostics)
        return [str(output_path)]

    is_csv = output_path.suffix.lower() == ".csv" or "csv" in lower_filter
    extension = ".csv" if is_csv else ".tsv"
    delimiter = "," if is_csv else "\t"
    base_name = output_path.stem
    for role in ("summary", "full", "diagnostics"):
        suffix = f"_{role}"
        if base_name.lower().endswith(suffix):
            base_name = base_name[: -len(suffix)]
            break

    exports = [("summary", summary), ("full", full)]
    if isinstance(diagnostics, TableExportData) and diagnostics.headers:
        exports.append(("diagnostics", diagnostics))

    written_paths: list[str] = []
    for role, data in exports:
        role_path = output_path.with_name(f"{base_name}_{role}{extension}")
        with role_path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh, delimiter=delimiter)
            writer.writerow(data.headers)
            writer.writerows(data.rows)
        written_paths.append(str(role_path))

    return written_paths


def _write_xlsx_bundle(
    output_path: Path,
    summary: TableExportData,
    full: TableExportData,
    diagnostics: TableExportData | None,
) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheets = [("关键结论", summary), ("完整明细", full)]
    if isinstance(diagnostics, TableExportData) and diagnostics.headers:
        sheets.append(("质量诊断", diagnostics))

    for index, (title, data) in enumerate(sheets):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = title
        _write_xlsx_row(worksheet, 1, data.headers)
        for row_index, row in enumerate(data.rows, start=2):
            _write_xlsx_row(worksheet, row_index, row)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(output_path)


def _write_xlsx_row(worksheet, row_index: int, values: tuple[Any, ...]) -> None:
    for column_index, value in enumerate(values, start=1):
        cell = worksheet.cell(row=row_index, column=column_index, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
