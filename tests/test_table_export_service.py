import csv
from dataclasses import FrozenInstanceError
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from polynexus.gui.table_export_service import (
    TableExportData,
    write_table_export,
    write_table_export_bundle,
)


def test_write_table_export_defaults_to_tsv(tmp_path):
    output_path = write_table_export(
        tmp_path / "results_table",
        ["Parameter", "Value"],
        [["L_nm", "12.0000"]],
    )

    assert output_path == str(tmp_path / "results_table.tsv")
    assert Path(output_path).read_text(encoding="utf-8").splitlines() == [
        "Parameter\tValue",
        "L_nm\t12.0000",
    ]


def test_write_table_export_uses_csv_when_filter_mentions_csv(tmp_path):
    output_path = write_table_export(
        tmp_path / "results_table_export",
        ["Parameter", "Value"],
        [["L_nm", "12.0000"]],
        selected_filter="CSV (*.csv)",
    )

    assert output_path == str(tmp_path / "results_table_export.csv")
    assert Path(output_path).read_text(encoding="utf-8").splitlines() == [
        "Parameter,Value",
        "L_nm,12.0000",
    ]


def test_table_export_data_is_immutable():
    data = TableExportData(headers=("Parameter", "Value"), rows=(("L_nm", 12.0),))

    with pytest.raises(FrozenInstanceError):
        data.headers = ("Changed",)


def test_write_table_export_bundle_writes_paired_csv_files(tmp_path):
    paths = write_table_export_bundle(
        tmp_path / "results.csv",
        summary=TableExportData(
            headers=("Metric", "Value"),
            rows=(("Mean", 12.5),),
        ),
        full=TableExportData(
            headers=("Index", "Value"),
            rows=((1, 12.5), (2, 13.75)),
        ),
        diagnostics=TableExportData(
            headers=("Level", "Message"),
            rows=(("warning", "Check input"),),
        ),
    )

    expected_paths = [
        tmp_path / "results_summary.csv",
        tmp_path / "results_full.csv",
        tmp_path / "results_diagnostics.csv",
    ]
    assert paths == [str(path) for path in expected_paths]
    with expected_paths[0].open(newline="", encoding="utf-8") as fh:
        assert list(csv.reader(fh)) == [["Metric", "Value"], ["Mean", "12.5"]]
    with expected_paths[1].open(newline="", encoding="utf-8") as fh:
        assert list(csv.reader(fh)) == [
            ["Index", "Value"],
            ["1", "12.5"],
            ["2", "13.75"],
        ]
    with expected_paths[2].open(newline="", encoding="utf-8") as fh:
        assert list(csv.reader(fh)) == [
            ["Level", "Message"],
            ["warning", "Check input"],
        ]


def test_write_table_export_bundle_defaults_to_tsv_and_omits_diagnostics(tmp_path):
    paths = write_table_export_bundle(
        tmp_path / "results",
        summary=TableExportData(headers=("Metric",), rows=(("Mean",),)),
        full=TableExportData(headers=("Index",), rows=((1,),)),
        diagnostics=TableExportData(headers=(), rows=()),
    )

    expected_paths = [
        tmp_path / "results_summary.tsv",
        tmp_path / "results_full.tsv",
    ]
    assert paths == [str(path) for path in expected_paths]
    assert expected_paths[0].read_text(encoding="utf-8").splitlines() == [
        "Metric",
        "Mean",
    ]
    assert expected_paths[1].read_text(encoding="utf-8").splitlines() == [
        "Index",
        "1",
    ]
    assert not (tmp_path / "results_diagnostics.tsv").exists()


def test_write_table_export_bundle_uses_filter_and_avoids_duplicate_role_suffix(tmp_path):
    paths = write_table_export_bundle(
        tmp_path / "results_summary",
        summary=TableExportData(headers=("Metric",), rows=()),
        full=TableExportData(headers=("Index",), rows=()),
        diagnostics=TableExportData(headers=("Issue",), rows=()),
        selected_filter="CSV (*.csv)",
    )

    assert paths == [
        str(tmp_path / "results_summary.csv"),
        str(tmp_path / "results_full.csv"),
        str(tmp_path / "results_diagnostics.csv"),
    ]


def test_write_table_export_bundle_writes_one_xlsx_with_typed_rows(tmp_path):
    output = tmp_path / "results.xlsx"

    paths = write_table_export_bundle(
        output,
        summary=TableExportData(
            headers=("Metric", "Value"),
            rows=(("Mean", 12.5),),
        ),
        full=TableExportData(
            headers=("Index", "Value"),
            rows=((1, 12.5), (2, 13.75)),
        ),
        diagnostics=TableExportData(
            headers=("Level", "Count"),
            rows=(("warning", 2),),
        ),
    )

    assert paths == [str(output)]
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["关键结论", "完整明细", "质量诊断"]
    assert list(workbook["关键结论"].values) == [
        ("Metric", "Value"),
        ("Mean", 12.5),
    ]
    assert list(workbook["完整明细"].values) == [
        ("Index", "Value"),
        (1, 12.5),
        (2, 13.75),
    ]
    assert workbook["完整明细"]["A2"].data_type == "n"
    for worksheet in workbook.worksheets:
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref == worksheet.dimensions


def test_write_table_export_bundle_keeps_formula_like_strings_as_literal_text(
    tmp_path,
):
    output = tmp_path / "literal_text.xlsx"
    exported_date = date(2026, 7, 10)
    write_table_export_bundle(
        output,
        summary=TableExportData(headers=("Metric",), rows=(("Mean",),)),
        full=TableExportData(
            headers=("Text", "Number", "Date", "Flag", "Empty"),
            rows=(("=1+1", 7, exported_date, True, None),),
        ),
        diagnostics=None,
    )

    workbook = load_workbook(output)
    values_workbook = load_workbook(output, data_only=True)
    try:
        for loaded_workbook in (workbook, values_workbook):
            row = loaded_workbook["完整明细"][2]
            assert row[0].value == "=1+1"
            assert row[1].value == 7
            loaded_date = row[2].value
            assert (
                loaded_date.date() if isinstance(loaded_date, datetime) else loaded_date
            ) == exported_date
            assert row[3].value is True
            assert row[4].value is None
        assert workbook["完整明细"]["A2"].data_type == "s"
    finally:
        workbook.close()
        values_workbook.close()


def test_write_table_export_bundle_uses_excel_filter_and_omits_diagnostics_sheet(
    tmp_path,
):
    paths = write_table_export_bundle(
        tmp_path / "results",
        summary=TableExportData(headers=("Metric",), rows=(("Mean",),)),
        full=TableExportData(headers=("Index",), rows=((1,),)),
        diagnostics=TableExportData(headers=(), rows=(("ignored",),)),
        selected_filter="Excel Workbook (*.xlsx)",
    )

    output = tmp_path / "results.xlsx"
    assert paths == [str(output)]
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["关键结论", "完整明细"]


@pytest.mark.parametrize("missing", ["summary", "full"])
def test_write_table_export_bundle_requires_both_main_headers_without_overwriting(
    tmp_path,
    missing,
):
    existing = tmp_path / "results_summary.csv"
    existing.write_text("keep me", encoding="utf-8")
    summary = TableExportData(headers=("Metric",), rows=(("Mean",),))
    full = TableExportData(headers=("Index",), rows=((1,),))
    if missing == "summary":
        summary = TableExportData(headers=(), rows=(("ignored",),))
    else:
        full = TableExportData(headers=(), rows=(("ignored",),))

    paths = write_table_export_bundle(
        tmp_path / "results.csv",
        summary=summary,
        full=full,
        diagnostics=TableExportData(headers=("Issue",), rows=(("ignored",),)),
    )

    assert paths == []
    assert existing.read_text(encoding="utf-8") == "keep me"
    assert not (tmp_path / "results_full.csv").exists()
    assert not (tmp_path / "results_diagnostics.csv").exists()
