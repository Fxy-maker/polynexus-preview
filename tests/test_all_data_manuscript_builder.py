from __future__ import annotations

import json
from pathlib import Path

import pytest


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def _minimal_package(tmp_path: Path) -> tuple[Path, Path]:
    project = tmp_path / "project"
    package = project / ".polynexus" / "evidence" / "pkg-v001"
    (project / "nmr").mkdir(parents=True)
    (project / "waxs").mkdir(parents=True)
    (project / "nmr" / "PA6-H.csv").write_text("0,1\n", encoding="utf-8")
    (project / "waxs" / "PA6.raw").write_text("raw", encoding="utf-8")
    run = {
        "run_id": "run-nmr",
        "status": "review_required",
        "analysis_run": {
            "recipe": {
                "artifacts": [
                    {
                        "technique": "nmr",
                        "path": "C:\\data\\nmr\\PA6-H.csv",
                        "format": "csv",
                    }
                ]
            },
            "evidence": {"supported_interpretations": ["nmr_spectrum:public_result_available"]},
        },
        "evidence_items": [{"evidence_id": "run-nmr:nmr_spectrum", "status": "review_required"}],
        "figures": [],
    }
    _write_json(package / "runs" / "run-nmr.json", run)
    _write_json(
        package / "manifest.json",
        {
            "package_id": "pkg",
            "version": 1,
            "status": "review_required",
            "run_ids": ["run-nmr"],
            "run_manifests": ["runs/run-nmr.json"],
            "evidence_count": 1,
            "figure_count": 0,
            "source_hashes": ["hash"],
            "limitations": ["nmr_spectrum:unqualified_export"],
            "package_hash": "package-hash",
            "techniques": {"nmr": {"run_ids": ["run-nmr"], "evidence_count": 1, "statuses": ["review_required"]}},
        },
    )
    return project, package


def _renderable_dsc_audit(*, noniso_rows: list[dict[str, str]], run_rows: list[dict[str, str]]) -> dict[str, object]:
    return {
        "package": {
            "package_id": "pkg",
            "version": 1,
            "status": "review_required",
            "run_count": len(run_rows),
            "evidence_count": 0,
            "figure_count": 0,
            "source_hash_count": 0,
            "techniques": {},
        },
        "project_root": "C:/project",
        "raw_file_count": 0,
        "sample_technique_rows": [],
        "run_rows": run_rows,
        "derived_tables": {
            "isothermal_dsc_pairwise": {"data": []},
            "isothermal_dsc_kinetics": {"data": []},
            "isothermal_dsc_kinetics_pa50": {"data": []},
            "nonisothermal_dsc_summary": {"data": noniso_rows},
        },
    }


def test_infer_sample_and_source_class_are_stable() -> None:
    from scripts.build_all_data_manuscript import infer_sample, classify_source

    assert infer_sample("solid-NMR/H/H-PA6TXT.csv") == "PA6"
    assert infer_sample("insu-FTIR/PA11-50/4033-SW-100.csv") == "PA11-50"
    assert classify_source(Path("DSC-升降升/PA6.opju"))["use"] == "audit_only"
    assert classify_source(Path("nmr/PA12-50-H.csv"))["technique"] == "NMR"


def test_build_audit_links_package_runs_and_keeps_review_boundary(tmp_path: Path) -> None:
    from scripts.build_all_data_manuscript import build_data_audit

    project, package = _minimal_package(tmp_path)
    audit = build_data_audit(project, package)
    rows = audit["source_rows"]
    nmr = next(row for row in rows if row["relative_path"] == "nmr/PA6-H.csv")
    assert nmr["sample"] == "PA6"
    assert nmr["package_run_ids"] == ["run-nmr"]
    assert nmr["review_status"] == "review_required"
    assert nmr["use"] == "audit_only"
    assert audit["package"]["run_count"] == 1


def test_builder_rejects_existing_output_without_force(tmp_path: Path) -> None:
    from scripts.build_all_data_manuscript import build_manuscript

    project, package = _minimal_package(tmp_path)
    output = project / "manuscript" / "even-nylon-crystallization-v002"
    output.mkdir(parents=True)
    (output / "keep.txt").write_text("do not overwrite", encoding="utf-8")
    with pytest.raises(FileExistsError):
        build_manuscript(project, package, output)


def test_docx_builder_applies_east_asia_font_to_base_styles(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document
    from docx.oxml.ns import qn
    from scripts.build_all_data_manuscript import _build_docx

    output = tmp_path / "manuscript.docx"
    assert _build_docx("# 标题\n\n正文", output, tmp_path) == "created"
    document = Document(output)
    for style_name in ("Normal", "Body Text"):
        style = document.styles[style_name]
        r_pr = style._element.rPr
        assert r_pr is not None
        assert r_pr.rFonts.get(qn("w:eastAsia")) == "SimSun"


def test_docx_tables_repeat_headers_and_keep_rows_together(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document
    from docx.oxml.ns import qn
    from scripts.build_all_data_manuscript import _add_docx_table

    document = Document()
    _add_docx_table(document, [["头", "值"], ["长内容", "第二行"]])

    table = document.tables[0]
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    assert header_tr_pr.find(qn("w:tblHeader")) is not None
    for row in table.rows:
        assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None


def test_builder_writes_all_core_outputs_for_a_new_directory(tmp_path: Path) -> None:
    from scripts.build_all_data_manuscript import build_manuscript

    project, package = _minimal_package(tmp_path)
    output = project / "manuscript" / "even-nylon-crystallization-v002"

    result = build_manuscript(project, package, output)

    assert Path(result["output_dir"]) == output
    for name in (
        "manuscript.md",
        "manuscript.json",
        "manuscript.docx",
        "preflight.json",
        "data_audit.csv",
        "data_audit.md",
        "run_audit.csv",
        "sample_technique_coverage.csv",
        "claim_evidence_literature_matrix.csv",
        "literature_matrix.csv",
        "figure_table_supplement_index.md",
        "supplementary_data_index.md",
        "review-status.md",
    ):
        assert (output / name).is_file(), name


def test_causal_boundary_allows_explicitly_negated_limits() -> None:
    from scripts.build_all_data_manuscript import _has_prohibited_causal_overclaim

    assert not _has_prohibited_causal_overclaim(
        "该结论不构成普适偶奇定律，也不是相同热力学过冷度下的本征速率比较。"
    )
    assert _has_prohibited_causal_overclaim("本研究证明了普适偶奇定律。")


def test_pairwise_tables_merge_pure_and_pa50_approved_kinetics() -> None:
    from scripts.build_all_data_manuscript import _pairwise_tables

    audit = {
        "derived_tables": {
            "isothermal_dsc_pairwise": {"data": []},
            "isothermal_dsc_kinetics": {"data": [{"sample": "PA6", "n": "2.0"}]},
            "isothermal_dsc_kinetics_pa50": {"data": [{"sample": "PA6-50", "n": "2.5"}]},
            "nonisothermal_dsc_summary": {"data": []},
        }
    }

    _, kinetics, _ = _pairwise_tables(audit)

    assert [(row["sample"], row["n"]) for row in kinetics] == [("PA6", "2.0"), ("PA6-50", "2.5")]


def test_source_strings_exclude_nested_metric_paths() -> None:
    from scripts.build_all_data_manuscript import _source_strings

    run = {
        "analysis_run": {"recipe": {"artifacts": []}},
        "metric_manifest": {"path": "result.nmr.assigned_peak_fraction"},
        "canonical": {"source_path": r"C:\\data\\nmr\\PA6-H.csv"},
        "diagnostic": {"path": "not-an-input"},
    }

    assert _source_strings(run) == [r"C:\\data\\nmr\\PA6-H.csv"]


def test_pairwise_range_summary_reads_all_four_approved_columns() -> None:
    from scripts.build_all_data_manuscript import _pairwise_range_summary

    summary = _pairwise_range_summary(
        [
            {
                "t_half_pure_min": "2.0",
                "t_half_PA_50_min": "1.0",
                "R_t_t_half_PA_50_over_pure": "0.5",
                "R_G_PA_50_over_pure": "2.0",
            },
            {
                "t_half_pure_min": "3.0",
                "t_half_PA_50_min": "1.5",
                "R_t_t_half_PA_50_over_pure": "0.6",
                "R_G_PA_50_over_pure": "1.7",
            },
        ]
    )

    assert summary == {
        "pure_t_half": "2–3",
        "pa50_t_half": "1–1.5",
        "time_ratio": "0.5–0.6",
        "rate_ratio": "1.7–2",
    }


def test_supplement_indexes_list_both_approved_kinetics_sources(tmp_path: Path) -> None:
    from scripts.build_all_data_manuscript import _write_figure_index, _write_supplement_index

    _write_figure_index(tmp_path)
    _write_supplement_index(tmp_path, {"run_rows": []})

    expected = "analysis_output/paper_data/dsc_kinetics.csv"
    assert expected in (tmp_path / "figure_table_supplement_index.md").read_text(encoding="utf-8")
    assert expected in (tmp_path / "supplementary_data_index.md").read_text(encoding="utf-8")


def test_temperature_parser_uses_terminal_or_hold_temperature_not_sample_label() -> None:
    from scripts.build_all_data_manuscript import _temperature_from_name

    assert _temperature_from_name(Path("PA11-JW-30.csv")) == "30"
    assert _temperature_from_name(Path("PA12-JW-30.csv")) == "30"
    assert _temperature_from_name(Path("4012-JW-30.csv")) == "30"
    assert _temperature_from_name(Path("PA11-220-for 1min.csv")) == "220"


def test_raw_match_does_not_cross_link_duplicate_basenames() -> None:
    from scripts.build_all_data_manuscript import _match_run_to_raw

    record = {
        "source_refs": [r"C:\\project\\DSC-升降升\\PA11.txt"],
        "source_basenames": ["pa11.txt"],
    }

    assert _match_run_to_raw(record, "DSC-升降升/PA11.txt")
    assert not _match_run_to_raw(record, "DSC-升降升/DSC数据/PA11.txt")


def test_generated_json_round_trips_all_shared_paper_contracts(tmp_path: Path) -> None:
    from polynexus.suite.paper_contracts import (
        CitationRequest,
        ClaimRecord,
        FigurePlan,
        FormulaRecord,
        ManuscriptSource,
    )
    from scripts.build_all_data_manuscript import build_manuscript

    project, package = _minimal_package(tmp_path)
    output = project / "manuscript" / "v002-r3"

    result = build_manuscript(project, package, output)
    payload = json.loads(Path(result["json"]).read_text(encoding="utf-8"))

    source = ManuscriptSource.from_dict(payload)
    claims = [ClaimRecord.from_dict(value) for value in payload["claims"]]
    figures = [FigurePlan.from_dict(value) for value in payload["figures"]]
    citations = [CitationRequest.from_dict(value) for value in payload["citations"]]
    formulas = [FormulaRecord.from_dict(value) for value in payload["formulas"]]

    assert source.claim_ids == tuple(value.claim_id for value in claims)
    assert source.figure_plan_ids == tuple(value.figure_id for value in figures)
    assert source.citation_ids == tuple(value.citation_id for value in citations)
    assert source.formula_ids == tuple(value.formula_id for value in formulas)
    assert source.to_dict()["projection"]["working_draft"]["title_zh"]
    assert {"svg", "json"} <= set(figures[0].outputs)


def test_citation_contract_uses_traceable_fallback_when_doi_is_missing() -> None:
    from scripts.build_all_data_manuscript import _citation_dicts

    citation = _citation_dicts([{"ref_id": "R99", "citation": "Local matrix record", "doi": ""}])[0]

    assert citation["locator"] == "local-matrix:R99"


def test_preflight_counts_both_approved_kinetics_tables() -> None:
    from scripts.build_all_data_manuscript import _custom_preflight

    audit = {
        "package": {"status": "review_required"},
        "raw_file_count": 1,
        "sample_technique_rows": [{"sample": sample} for sample in ("PA6", "PA6-50", "PA11", "PA11-50", "PA12", "PA12-50")],
        "derived_tables": {
            "isothermal_dsc_pairwise": {"rows": 15},
            "isothermal_dsc_kinetics": {"rows": 15},
            "isothermal_dsc_kinetics_pa50": {"rows": 15},
            "nonisothermal_dsc_summary": {"rows": 6},
        },
    }

    preflight = _custom_preflight({"source_id": "source"}, "不构成普适偶奇定律。", audit, "created")

    assert preflight["checks"]["dsc_kinetics"] == {"status": "pass", "rows": 30}


def test_main_figures_use_current_pairwise_assets() -> None:
    from scripts.build_all_data_manuscript import FIGURES

    figures = {figure["id"]: figure for figure in FIGURES}

    assert figures["F1"]["path"].endswith("Fig2_nonisothermal_dsc_thermal_curves.png")
    assert figures["F4"]["path"].endswith("Fig8_pure_pa_vs_pa50_kinetics.png")
    assert figures["F6"]["path"].endswith("Fig9_pure_pa_vs_pa50_scattering.png")


def test_figure_metadata_keeps_technique_and_attachment_boundaries() -> None:
    from scripts.build_all_data_manuscript import FIGURES

    figures = {figure["id"]: figure for figure in FIGURES}

    f2 = figures["F2"]
    assert "等温" in f2["title"]
    assert "非等温" not in f2["title"]
    assert "analysis_output/control_dsc/dsc_candidate_curves.csv" in f2["source"]
    assert "analysis_output/paper_data/dsc_candidate_curves.csv" in f2["source"]
    assert "candidate_curves" in f2["boundary"]

    s2 = figures["S2"]
    assert "12 溶液" in s2["title"]
    assert "14 固体" in s2["title"]
    assert "仅审计" in s2["title"]
    assert "12 solution NMR runs" in s2["source"]
    assert "14 solid NMR runs" in s2["source"]
    assert "仅审计" in s2["boundary"]

    f6 = figures["F6"]
    assert "外部派生附件" in f6["source"]
    assert "analysis_output/paper_data/scattering_profiles.csv" in f6["source"]
    assert "证据包之外" in f6["boundary"]


def test_interpretation_claims_link_both_members_of_each_dsc_pair() -> None:
    from scripts.build_all_data_manuscript import build_claim_matrix

    run_rows = []
    for hard_segment in ("PA6", "PA11", "PA12"):
        for sample in (hard_segment, f"{hard_segment}-50"):
            for mode in ("isothermal_dsc", "nonisothermal_dsc"):
                run_rows.append(
                    {
                        "technique": "dsc",
                        "samples": [sample],
                        "source_role": mode,
                        "package_evidence_ids": [f"{sample}:{mode}"],
                    }
                )

    claims = {item["claim_id"]: item for item in build_claim_matrix({"package": {}, "run_rows": run_rows})}
    expected_paths = (
        "analysis_output/paper_figures_pairwise/nonisothermal_dsc_thermal_summary.csv + "
        "analysis_output/control_dsc/dsc_pure_vs_pa50_pairwise.csv"
    )
    for claim_id, hard_segment in (("C07", "PA6"), ("C08", "PA12"), ("C09", "PA11")):
        claim = claims[claim_id]
        assert claim["source_paths"] == expected_paths
        evidence = set(claim["package_evidence_ids"].split(";"))
        assert evidence == {
            f"{hard_segment}:isothermal_dsc",
            f"{hard_segment}:nonisothermal_dsc",
            f"{hard_segment}-50:isothermal_dsc",
            f"{hard_segment}-50:nonisothermal_dsc",
        }


def test_f2_marker_is_rendered_with_isothermal_dsc_section() -> None:
    from scripts.build_all_data_manuscript import _render_manuscript_markdown

    markdown = _render_manuscript_markdown(
        _renderable_dsc_audit(noniso_rows=[], run_rows=[]),
        [],
        [],
    )
    marker = "<!--FIGURE:F2|"
    assert markdown.index("### 3.3 等温 DSC") < markdown.index(marker) < markdown.index("### 3.4 表观 Avrami")


def test_rendered_nonisothermal_dsc_prefers_approved_delta_t_field() -> None:
    from scripts.build_all_data_manuscript import _render_manuscript_markdown

    audit = _renderable_dsc_audit(
        noniso_rows=[
            {
                "sample": "PA11",
                "Tm_C": "190.0",
                "Tc_C": "120.0",
                "DeltaT_C": "70.0",
                "DeltaHm_J_g": "1.0",
                "DeltaHc_J_g": "1.0",
            },
            {
                "sample": "PA11-50",
                "Tm_C": "180.0",
                "Tc_C": "110.2",
                "DeltaT_C": "69.9",
                "DeltaHm_J_g": "1.0",
                "DeltaHc_J_g": "1.0",
            }
        ],
        run_rows=[],
    )

    markdown = _render_manuscript_markdown(audit, [], [])

    assert "| PA11-50 | 180 | 110.2 | 69.9 | 1 | 1 |" in markdown
    assert "PA11 由 70 °C 变为 69.9 °C" in markdown
    assert "69.8" not in markdown


def test_rendered_dsc_summary_uses_actual_source_role_counts() -> None:
    from scripts.build_all_data_manuscript import _render_manuscript_markdown

    audit = _renderable_dsc_audit(
        noniso_rows=[],
        run_rows=[
            *[{"technique": "dsc", "source_role": "isothermal_dsc"} for _ in range(8)],
            *[{"technique": "dsc", "source_role": "nonisothermal_dsc"} for _ in range(10)],
            {"technique": "ir", "source_role": "isothermal_dsc"},
            {"technique": "dsc", "source_role": "unknown"},
        ],
    )

    markdown = _render_manuscript_markdown(audit, [], [])

    assert "DSC 的 18 个包内 runs（8 个等温来源和 10 个非等温文本来源）" in markdown

    small_audit = _renderable_dsc_audit(
        noniso_rows=[],
        run_rows=[
            *[{"technique": "dsc", "source_role": "isothermal_dsc"} for _ in range(2)],
            {"technique": "dsc", "source_role": "nonisothermal_dsc"},
        ],
    )

    small_markdown = _render_manuscript_markdown(small_audit, [], [])

    assert "DSC 的 3 个包内 runs（2 个等温来源和 1 个非等温文本来源）" in small_markdown


def test_claim_matrix_uses_nonisothermal_tm_minus_tc_definition() -> None:
    from scripts.build_all_data_manuscript import build_claim_matrix

    claims = build_claim_matrix({"package": {}, "run_rows": []})
    claim = next(item for item in claims if item["claim_id"] == "C07")

    assert "Tm−Tc" in claim["claim"]
    assert "Tm−Tiso" not in claim["claim"]
